# update_tts/run.py
# Entrypoint: python -m update_tts.run "duong_dan_file.xlsx" [--dry-run] [--observe]
#
# --dry-run: điền thử form nhưng bấm HỦY thay vì XÁC NHẬN thật.
# --observe: CHẠY CHẬM LẠI, dừng 5 giây sau mỗi bước để quan sát bằng mắt.
#            Tự động bật kèm --dry-run (không thể quan sát mà lại ghi dữ liệu thật).

import sys
import time
from playwright.sync_api import sync_playwright

# Fix UnicodeEncodeError trên terminal Windows (cp1252 không in được emoji)
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

from . import config
from . import browser_utils
from . import excel_reader
from . import ticket_actions
from ai_cache import clear_cache
from pathlib import Path

OBSERVE_STEP_DELAY_MS = 5000
def get_latest_result_file():
    project_root = Path(__file__).resolve().parent.parent
    result_dir = project_root / "result"

    if not result_dir.exists():
        return None

    excel_files = [
        path for path in result_dir.glob("*.xlsx")
        if not path.name.startswith("~$")
    ]

    if not excel_files:
        return None

    return max(excel_files, key=lambda path: path.stat().st_mtime)

def run_update_tts(excel_path=None, dry_run=False, observe=False):
    """
    Hàm thực thi cập nhật và đóng phiếu TTS tự động từ file Excel kết quả.
    Có thể gọi từ main.py hoặc chạy độc lập.
    """
    if excel_path is None:
        excel_path = get_latest_result_file()

    if excel_path is None or not Path(excel_path).exists():
        print("❌ Không tìm thấy file Excel nào trong thư mục result/.")
        return False

    excel_path = Path(excel_path)
    print(f"\n📄 [UPDATE TTS] File Excel đang dùng: {excel_path}")

    step_delay_ms = OBSERVE_STEP_DELAY_MS if observe else 0

    if observe:
        print(f"🐢 CHẾ ĐỘ QUAN SÁT: dừng {OBSERVE_STEP_DELAY_MS/1000:.0f} giây sau mỗi bước để bạn theo dõi.")
    if dry_run:
        print("🧪 CHẠY Ở CHẾ ĐỘ DRY-RUN: sẽ điền thử form nhưng KHÔNG bấm XÁC NHẬN thật (bấm HỦY thay thế).\n")

    records = excel_reader.read_excel_results(excel_path)
    auto_close_records = [
        rec for rec in records
        if excel_reader.is_level_1_auto_close_candidate(rec)
    ]
    excluded_count = len(records) - len(auto_close_records)

    print(
        f"🎯 {len(auto_close_records)} phiếu đủ điều kiện tự động cập nhật + đóng "
        f"(bỏ qua {excluded_count} phiếu chưa đủ điều kiện)."
    )

    records = auto_close_records
    if not records:
        print("⚠️ Không có dữ liệu nào trong file Excel đủ điều kiện để cập nhật.")
        return True

    with sync_playwright() as p:
        try:
            browser = browser_utils.connect_to_chrome(p)
            print("✅ Đã kết nối vào Chrome debug thành công!")
        except Exception as ex:
            print(f"❌ Không kết nối được Chrome Debugging Port 9222: {ex}")
            return False

        context = browser_utils.get_context(browser)
        page = browser_utils.find_tts_page(context)

        if page is None:
            print("❌ Không tìm thấy tab TTS đang mở. Vui lòng kiểm tra lại tab TTS.")
            browser.close()
            return False

        print(f"📄 Đang dùng tab TTS: {page.url}")

        success_count = 0
        skip_count = 0
        fail_count = 0

        for rec in records:
            print(f"\n🔧 Đang xử lý SĐT {rec['phone']} (trạng thái: {rec['status']})...")
            result = ticket_actions.close_ticket(
                page,
                rec["phone"],
                rec["status"],
                rec["comment"],       # NỘI DUNG PHÂN TÍCH KỸ THUẬT
                rec["action_plan"],   # HƯỚNG XỬ LÝ KHUYÊN DÙNG
                dry_run=dry_run,
                step_delay_ms=step_delay_ms
            )
            if result is True:
                success_count += 1
                rec["ticket_status"] = "Đã đóng"
            elif result is None:
                skip_count += 1
            else:
                fail_count += 1
            time.sleep(1)  # đệm nhẹ giữa các lần thao tác

        # Đồng bộ trạng thái 'Đã đóng' vào Cột J trong file Excel kết quả & Database SQLite
        if success_count > 0 and not dry_run:
            try:
                import openpyxl
                from openpyxl.styles import Font, PatternFill, Alignment
                wb_sync = openpyxl.load_workbook(excel_path)
                ws_sync = wb_sync.active
                closed_phones = {r["phone"] for r in records if r.get("ticket_status") == "Đã đóng"}
                for r_idx in range(config.EXCEL_DATA_START_ROW, ws_sync.max_row + 1):
                    p_num = str(ws_sync.cell(row=r_idx, column=config.EXCEL_COL_PHONE + 1).value or "").strip()
                    if p_num in closed_phones:
                        st_c = ws_sync.cell(row=r_idx, column=config.EXCEL_COL_TICKET_STATUS + 1)
                        st_c.value = "Đã đóng"
                        st_c.font = Font(name="Segoe UI", size=10, bold=True, color="006100")
                        st_c.fill = PatternFill(start_color="C6EFCE", fill_type="solid")
                        st_c.alignment = Alignment(horizontal="center", vertical="center")
                wb_sync.save(excel_path)
            except Exception as ex_sync:
                print(f"⚠️ Không thể cập nhật trạng thái Excel: {ex_sync}")

            try:
                from db_manager import update_ticket_field
                for r in records:
                    if r.get("ticket_status") == "Đã đóng":
                        update_ticket_field(r["phone"], "ticket_status", "Đã đóng")
            except Exception as ex_db:
                print(f"⚠️ Không thể cập nhật trạng thái Database: {ex_db}")

        print("\n===================================")
        print(f"✅ Thành công: {success_count}")
        print(f"⏭️  Bỏ qua (chưa có mapping / phiếu mở lại): {skip_count}")
        print(f"❌ Thất bại: {fail_count}")
        if dry_run:
            print("🧪 (Đây là kết quả DRY-RUN, chưa có gì được ghi thật lên TTS)")
        else:
            if success_count > 0:
                try:
                    clear_cache()
                except Exception:
                    pass
        print("===================================")

        browser.close()
        return True


def close_single_ticket_from_db(phone, incident_time=None, dry_run=False, observe=False):
    """
    Đóng 1 phiếu cụ thể từ database trên hệ thống TTS qua Chrome Debugging (Port 9222).
    Trả về (success: bool, message: str)
    """
    from db_manager import get_all_tickets, update_ticket_field

    tickets = get_all_tickets()
    target_ticket = None
    clean_p = "".join(filter(str.isdigit, str(phone or "")))
    
    # 1. Tìm chính xác theo phone và incident_time (nếu có)
    for t in tickets:
        t_p = "".join(filter(str.isdigit, str(t.get("phone", ""))))
        if t_p == clean_p or (len(t_p) >= 9 and len(clean_p) >= 9 and (t_p.endswith(clean_p[-9:]) or clean_p.endswith(t_p[-9:]))):
            if incident_time and incident_time != "--" and str(t.get("incident_time", "")).strip():
                if str(t.get("incident_time", "")).strip() == str(incident_time).strip():
                    target_ticket = t
                    break
            else:
                target_ticket = t
                break

    # 2. Dự phòng: Tìm theo phone nếu chưa khớp incident_time
    if not target_ticket:
        for t in tickets:
            t_p = "".join(filter(str.isdigit, str(t.get("phone", ""))))
            if t_p == clean_p or (len(t_p) >= 9 and len(clean_p) >= 9 and (t_p.endswith(clean_p[-9:]) or clean_p.endswith(t_p[-9:]))):
                target_ticket = t
                break

    if not target_ticket:
        return False, f"Không tìm thấy phiếu của SĐT {phone} trong cơ sở dữ liệu."

    status = target_ticket.get("status", "")
    comment = target_ticket.get("comment", "")
    action_plan = target_ticket.get("action_plan", "")

    step_delay_ms = OBSERVE_STEP_DELAY_MS if observe else 0

    with sync_playwright() as p:
        try:
            browser = browser_utils.connect_to_chrome(p)
        except Exception as ex:
            return False, f"Không kết nối được Chrome Debugging Port 9222: {ex}"

        context = browser_utils.get_context(browser)
        page = browser_utils.find_tts_page(context)

        if page is None:
            browser.close()
            return False, "Không tìm thấy tab TTS đang mở trên Chrome."

        result = ticket_actions.close_ticket(
            page,
            phone,
            status,
            comment,
            action_plan,
            dry_run=dry_run,
            step_delay_ms=step_delay_ms
        )

        browser.close()

        if result is True:
            if not dry_run:
                update_ticket_field(phone, "ticket_status", "Đã đóng", incident_time=incident_time)
                try:
                    clear_cache()
                except Exception:
                    pass
            return True, f"Đã đóng thành công phiếu cho SĐT {phone} trên TTS."
        elif result is None:
            return False, f"Bỏ qua: Phiếu SĐT {phone} có thể là 'Phiếu mở lại' hoặc trạng thái '{status}' chưa có mapping nguyên nhân."
        else:
            return False, f"Thao tác đóng phiếu thất bại cho SĐT {phone}. Vui lòng kiểm tra lại tab TTS."


def main():
    args = sys.argv[1:]
    observe = "--observe" in args
    dry_run = ("--dry-run" in args) or observe

    excel_args = [arg for arg in args if not arg.startswith("--")]
    excel_path = Path(excel_args[0]) if excel_args else None

    run_update_tts(excel_path=excel_path, dry_run=dry_run, observe=observe)


if __name__ == "__main__":
    main()
