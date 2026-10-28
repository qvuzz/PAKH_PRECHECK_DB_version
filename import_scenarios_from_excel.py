# import_scenarios_from_excel.py
# Cong cu doc lai file Excel KICH_BAN_TIEN_KIEM_PRECHECK.xlsx sau khi KTV da sua cau tu
# va dong bo cap nhat vao cac file JSON (sms_diagnostic_scenarios.json, diagnostic_scenarios_sms.json, diagnostic_scenarios.json)

import os
import json
import openpyxl

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
EXCEL_FILE = os.path.join(BASE_DIR, "KICH_BAN_TIEN_KIEM_PRECHECK.xlsx")

def import_scenarios(dry_run=True, sheet_target="Kich_Ban_Tin_Nhan_SMS"):
    if not os.path.exists(EXCEL_FILE):
        print(f"[ERR] Khong tim thay file Excel: {EXCEL_FILE}")
        return

    wb = openpyxl.load_workbook(EXCEL_FILE, data_only=True)
    
    # 1. Doc Sheet SMS
    if sheet_target in ("all", "Kich_Ban_Tin_Nhan_SMS") and "Kich_Ban_Tin_Nhan_SMS" in wb.sheetnames:
        ws_sms = wb["Kich_Ban_Tin_Nhan_SMS"]
        sms_dict = {}
        for row in range(2, ws_sms.max_row + 1):
            sc_id = ws_sms.cell(row=row, column=2).value
            if not sc_id:
                continue
            sc_id = str(sc_id).strip()
            sms_dict[sc_id] = {
                "id": sc_id,
                "name": str(ws_sms.cell(row=row, column=3).value or "").strip(),
                "conclusion": str(ws_sms.cell(row=row, column=4).value or "").strip(),
                "suggested_cause_code": str(ws_sms.cell(row=row, column=5).value or "").strip(),
                "suggested_cause_name": str(ws_sms.cell(row=row, column=6).value or "").strip(),
                "technical_analysis": str(ws_sms.cell(row=row, column=7).value or "").strip(),
                "action_plan": str(ws_sms.cell(row=row, column=8).value or "").strip(),
                "color": str(ws_sms.cell(row=row, column=9).value or "").strip() or "green"
            }
        
        # So sanh voi du lieu hien tai de tim diffs
        ref_path = os.path.join(BASE_DIR, "sms_diagnostic_scenarios.json")
        ref_scenarios = {}
        if os.path.exists(ref_path):
            with open(ref_path, "r", encoding="utf-8") as f:
                ref_data = json.load(f)
                r_list = ref_data.get("SMS_SCENARIOS") or ref_data.get("SCENARIOS") or []
                ref_scenarios = {x["id"]: x for x in r_list}

        fields_to_check = [
            ("name", "Tên kịch bản"),
            ("conclusion", "Kết luận"),
            ("suggested_cause_code", "Mã lỗi OneOSS"),
            ("suggested_cause_name", "Tên mã lỗi OneOSS"),
            ("technical_analysis", "Phân tích kỹ thuật"),
            ("action_plan", "Hướng xử lý"),
            ("color", "Màu sắc")
        ]

        total_changed = 0
        diff_summary = []
        for sc_id, new_item in sms_dict.items():
            old_item = ref_scenarios.get(sc_id, {})
            item_changes = []
            for f_key, f_label in fields_to_check:
                old_v = str(old_item.get(f_key, "")).strip()
                new_v = str(new_item.get(f_key, "")).strip()
                if old_v != new_v:
                    item_changes.append((f_key, f_label, old_v, new_v))
            if item_changes:
                total_changed += 1
                diff_summary.append((sc_id, new_item.get("name") or old_item.get("name"), item_changes))

        # Ghi log diff ra file utf-8 de xem chi tiet an toan
        diff_log_path = os.path.join(BASE_DIR, "sms_scenarios_diff.log")
        with open(diff_log_path, "w", encoding="utf-8") as lf:
            lf.write(f"=== TỔNG HỢP {total_changed} KỊCH BẢN SMS ĐƯỢC CHỈNH SỬA TỪ EXCEL ===\n\n")
            for sc_id, s_name, ch_list in diff_summary:
                lf.write(f"[{sc_id}] - {s_name}\n")
                for f_key, f_label, old_v, new_v in ch_list:
                    lf.write(f"  * {f_label} ({f_key}):\n")
                    lf.write(f"    - TRƯỚC: {old_v}\n")
                    lf.write(f"    + SAU:   {new_v}\n")
                lf.write("\n")

        print(f"[DIFF CHECK] Phat hien {total_changed} kich ban SMS co thay doi noi dung.")
        print(f"[DIFF LOG] Chi tiet luu tai: {diff_log_path}")

        # Cap nhat vao sms_diagnostic_scenarios.json & diagnostic_scenarios_sms.json
        for f_name in ["sms_diagnostic_scenarios.json", "diagnostic_scenarios_sms.json"]:
            f_path = os.path.join(BASE_DIR, f_name)
            if os.path.exists(f_path):
                with open(f_path, "r", encoding="utf-8") as f:
                    data = json.load(f)
                
                # Check SMS_SCENARIOS hoac SCENARIOS
                for key in ["SMS_SCENARIOS", "SCENARIOS"]:
                    if key in data and isinstance(data[key], list):
                        for item in data[key]:
                            item_id = item.get("id")
                            if item_id in sms_dict:
                                new_val = sms_dict[item_id]
                                if new_val["name"]: item["name"] = new_val["name"]
                                if new_val["conclusion"]: item["conclusion"] = new_val["conclusion"]
                                if new_val["suggested_cause_code"]: item["suggested_cause_code"] = new_val["suggested_cause_code"]
                                if new_val["suggested_cause_name"]: item["suggested_cause_name"] = new_val["suggested_cause_name"]
                                if new_val["technical_analysis"]: item["technical_analysis"] = new_val["technical_analysis"]
                                if new_val["action_plan"]: item["action_plan"] = new_val["action_plan"]
                                if new_val["color"]: item["color"] = new_val["color"]
                
                if not dry_run:
                    with open(f_path, "w", encoding="utf-8") as f:
                        json.dump(data, f, ensure_ascii=False, indent=4)
                    print(f"[APPLY] Da cap nhat thanh cong: {f_name}")
                else:
                    print(f"[DRY-RUN] San sang cap nhat vao: {f_name}")

    # 2. Doc Sheet Data Mobile (neu chon all hoac Kich_Ban_Data_Mobile)
    if sheet_target in ("all", "Kich_Ban_Data_Mobile") and "Kich_Ban_Data_Mobile" in wb.sheetnames:
        ws_data = wb["Kich_Ban_Data_Mobile"]
        data_dict = {}
        for row in range(2, ws_data.max_row + 1):
            sc_id = ws_data.cell(row=row, column=2).value
            if not sc_id:
                continue
            sc_id = str(sc_id).strip()
            data_dict[sc_id] = {
                "id": sc_id,
                "name": str(ws_data.cell(row=row, column=3).value or "").strip(),
                "conclusion": str(ws_data.cell(row=row, column=4).value or "").strip(),
                "technical_analysis": str(ws_data.cell(row=row, column=5).value or "").strip(),
                "action_plan": str(ws_data.cell(row=row, column=6).value or "").strip(),
                "color": str(ws_data.cell(row=row, column=7).value or "").strip() or "F2F2F2"
            }
        
        f_path = os.path.join(BASE_DIR, "diagnostic_scenarios.json")
        if os.path.exists(f_path):
            with open(f_path, "r", encoding="utf-8") as f:
                d_json = json.load(f)
            
            if "SCENARIOS" in d_json and isinstance(d_json["SCENARIOS"], list):
                for item in d_json["SCENARIOS"]:
                    item_id = item.get("id")
                    if item_id in data_dict:
                        new_val = data_dict[item_id]
                        if new_val["name"]: item["name"] = new_val["name"]
                        if new_val["conclusion"]: item["conclusion"] = new_val["conclusion"]
                        if new_val["technical_analysis"]: item["technical_analysis"] = new_val["technical_analysis"]
                        if new_val["action_plan"]: item["action_plan"] = new_val["action_plan"]
                        if new_val["color"]: item["color"] = new_val["color"]
            
            if not dry_run:
                with open(f_path, "w", encoding="utf-8") as f:
                    json.dump(d_json, f, ensure_ascii=False, indent=4)
                print(f"[APPLY] Da cap nhat thanh cong: diagnostic_scenarios.json")
            else:
                print(f"[DRY-RUN] San sang cap nhat vao: diagnostic_scenarios.json")

if __name__ == "__main__":
    import sys
    dry_mode = "--apply" not in sys.argv
    target = "Kich_Ban_Tin_Nhan_SMS"
    if "--all" in sys.argv:
        target = "all"
    elif "--data" in sys.argv:
        target = "Kich_Ban_Data_Mobile"
    import_scenarios(dry_run=dry_mode, sheet_target=target)
