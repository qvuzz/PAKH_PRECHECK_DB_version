# -*- coding: utf-8 -*-
"""
Script tạo Slide trình chiếu PowerPoint (.pptx) chuẩn phong cách nhận diện VNPT
PHIÊN BẢN TOÀN DIỆN 18 SLIDE - BAO QUÁT TOÀN BỘ 6 PHÂN HỆ VIỄN THÔNG:
1. Mobile Internet (Data 4G/5G)
2. Cuộc gọi & Thoại (Voice Calls, VoLTE)
3. Chặn gọi ngoại mạng / Chặn Spam (Nghị định 91, Cam kết CCOS)
4. Tin nhắn (SMS) & Cổng tra cứu SMSC CDR Gateway
5. CVQT (Roaming) & SIM / eSIM
6. Quản trị phân vùng KTV & Cấu hình SOC
"""

import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

ARTIFACT_DIR = r"C:\Users\lequa\.gemini\antigravity-ide\brain\233538e9-b8ef-4c40-b0c3-e80e1868d2c3"
IMG_DIR = os.path.join(ARTIFACT_DIR, "manual_crops")

# Các ảnh chụp từ hệ thống
IMG_HEADER = os.path.join(IMG_DIR, "ui_header.png")
IMG_SIDEBAR = os.path.join(IMG_DIR, "ui_sidebar.png")
IMG_ACTION_TABS = os.path.join(IMG_DIR, "ui_action_tabs.png")
IMG_EXPANDED = os.path.join(IMG_DIR, "ui_expanded_all.png")
IMG_CARD2 = os.path.join(IMG_DIR, "ui_card2_profile.png")
IMG_BTNS = os.path.join(IMG_DIR, "ui_action_buttons.png")
IMG_DROPDOWN = os.path.join(IMG_DIR, "ui_cause_dropdown.png")
IMG_WARD = os.path.join(IMG_DIR, "ui_ward_detect.png")

# Các ảnh phân hệ mới
IMG_VOICE = os.path.join(ARTIFACT_DIR, "module_cuoc_goi_1791814562393.png")
IMG_SPAM = os.path.join(ARTIFACT_DIR, "module_chan_spam_1791814585623.png")
IMG_SMS = os.path.join(ARTIFACT_DIR, "module_tin_nhan_1791814616208.png")
IMG_SMS_MODAL = os.path.join(ARTIFACT_DIR, "smsc_cdr_modal_1791814637383.png")
IMG_KTV_MGT = os.path.join(ARTIFACT_DIR, "quan_ly_ktv_1791814684329.png")

# Màu sắc VNPT
C_VNPT_BLUE = RGBColor(0, 82, 165)
C_VNPT_ACCENT = RGBColor(2, 132, 199)
C_DARK = RGBColor(15, 23, 42)
C_MUTED = RGBColor(100, 116, 139)
C_WHITE = RGBColor(255, 255, 255)
C_CARD_BORDER = RGBColor(226, 232, 240)
C_WARN = RGBColor(217, 119, 6)
C_SUCCESS = RGBColor(22, 163, 74)
C_DANGER = RGBColor(220, 38, 38)

def create_full_slide_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    def add_header(slide, title_text, category="VNPT TTS PRECHECK - HƯỚNG DẪN SỬ DỤNG TOÀN DIỆN"):
        top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.1))
        top_bar.fill.solid()
        top_bar.fill.fore_color.rgb = C_VNPT_BLUE
        top_bar.line.color.rgb = C_VNPT_BLUE

        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.25), Inches(11.7), Inches(0.9))
        tf = title_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p_cat = tf.paragraphs[0]
        p_cat.text = category.upper()
        p_cat.font.name = "Arial"
        p_cat.font.size = Pt(10)
        p_cat.font.bold = True
        p_cat.font.color.rgb = C_VNPT_ACCENT

        p_title = tf.add_paragraph()
        p_title.text = title_text
        p_title.font.name = "Arial"
        p_title.font.size = Pt(20)
        p_title.font.bold = True
        p_title.font.color.rgb = C_VNPT_BLUE

    def add_card(slide, left, top, width, height, bg_color=C_WHITE, border_color=C_CARD_BORDER):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        card.line.color.rgb = border_color
        card.line.width = Pt(1)
        return card

    # =========================================================================
    # SLIDE 1: TRANG BÌA (TITLE)
    # =========================================================================
    slide1 = prs.slides.add_slide(blank_layout)
    bg1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = C_VNPT_BLUE
    bg1.line.color.rgb = C_VNPT_BLUE

    card_title = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(1.2), Inches(10.933), Inches(5.1))
    card_title.fill.solid()
    card_title.fill.fore_color.rgb = C_WHITE
    card_title.line.color.rgb = C_WHITE

    tb1 = slide1.shapes.add_textbox(Inches(1.8), Inches(1.5), Inches(9.7), Inches(4.5))
    tf1 = tb1.text_frame
    tf1.word_wrap = True

    p = tf1.paragraphs[0]
    p.text = "TẬP ĐOÀN BƯU CHÍNH VIỄN THÔNG VIỆT NAM (VNPT) - KHỐI KỸ THUẬT"
    p.font.name = "Arial"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_VNPT_ACCENT

    p = tf1.add_paragraph()
    p.space_before = Pt(12)
    p.text = "HƯỚNG DẪN SỬ DỤNG HỆ THỐNG\nVNPT TTS PRECHECK"
    p.font.name = "Arial"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = C_VNPT_BLUE

    p = tf1.add_paragraph()
    p.space_before = Pt(8)
    p.text = "CẨM NANG TOÀN DIỆN 6 PHÂN HỆ: DATA, CUỘC GỌI, SPAM, SMS, CVQT, SIM & QUẢN TRỊ KTV"
    p.font.name = "Arial"
    p.font.size = Pt(12.5)
    p.font.bold = True
    p.font.color.rgb = C_DANGER

    p = tf1.add_paragraph()
    p.space_before = Pt(16)
    p.text = "• Hướng dẫn chi tiết từng chức năng, giao diện & nút bấm cho người mới bắt đầu (Newbie)\n• Tiền kiểm 4 hệ thống: Core SAPC, BTools CDP, Sóng CEM và SMSC CDR Gateway\n• Chuẩn hóa đóng phiếu OneOSS 5.1: Bắt buộc địa bàn 2 cấp theo 34 Tỉnh/TP mới sáp nhập"
    p.font.name = "Arial"
    p.font.size = Pt(10.5)
    p.font.color.rgb = C_DARK

    p = tf1.add_paragraph()
    p.space_before = Pt(20)
    p.text = "Cổng truy cập: http://10.155.59.158:1234  |  Phiên bản 2.6 (Tháng 10/2026)"
    p.font.name = "Arial"
    p.font.size = Pt(10)
    p.font.italic = True
    p.font.color.rgb = C_MUTED

    # =========================================================================
    # SLIDE 2: TỔNG QUAN 6 PHÂN HỆ NGHIỆP VỤ
    # =========================================================================
    slide2 = prs.slides.add_slide(blank_layout)
    add_header(slide2, "1. TỔNG QUAN 6 PHÂN HỆ NGHIỆP VỤ VIỄN THÔNG")

    mods_overview = [
        ("MOBILE INTERNET (DATA)", "Mạng 3G/4G/5G, bóp băng thông, OCS ODA_BIG_OCS, CEM sóng trạm."),
        ("CUỘC GỌI & THOẠI", "Rớt cuộc gọi, không gọi đi/đến, cờ khóa cước NAM, cấu hình VoLTE/HSS."),
        ("CHẶN SPAM NGOẠI MẠNG", "Bị Viettel/MobiFone chặn liên mạng NĐ 91, đối soát biên bản cam kết CCOS."),
        ("TIN NHẮN SMS & SMSC CDR", "Không nhận được OTP ngân hàng, đối soát mã lỗi cổng SMSC Gateway thời gian thực."),
        ("CVQT (ROAMING) & SIM", "Khách đi nước ngoài mất sóng, chuyển vùng quốc tế, đổi SIM 4G, quét eSIM."),
        ("QUẢN TRỊ KTV & SOC", "Phân quyền KTV 3 Miền (MB, MT, MN), kết nối TTS OneOSS, nhật ký Live Log.")
    ]
    for i, (m_title, m_desc) in enumerate(mods_overview):
        row = i // 3
        col = i % 3
        left = Inches(0.8 + col * 4.0)
        top = Inches(1.4 + row * 2.8)
        add_card(slide2, left, top, Inches(3.7), Inches(2.5))

        tb = slide2.shapes.add_textbox(left + Inches(0.2), top + Inches(0.2), Inches(3.3), Inches(2.1))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = f"PHÂN HỆ 0{i+1}"
        p.font.name = "Arial"
        p.font.size = Pt(9.5)
        p.font.bold = True
        p.font.color.rgb = C_VNPT_ACCENT

        p = tf.add_paragraph()
        p.space_before = Pt(4)
        p.text = m_title
        p.font.name = "Arial"
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = C_VNPT_BLUE

        p = tf.add_paragraph()
        p.space_before = Pt(6)
        p.text = m_desc
        p.font.name = "Arial"
        p.font.size = Pt(10)
        p.font.color.rgb = C_DARK

    # =========================================================================
    # SLIDE 3: GIAO DIỆN CHÍNH & THANH HEADER
    # =========================================================================
    slide3 = prs.slides.add_slide(blank_layout)
    add_header(slide3, "2. GIAO DIỆN TỔNG QUAN & THANH HEADER ĐIỀU KHIỂN")

    if os.path.exists(IMG_HEADER):
        slide3.shapes.add_picture(IMG_HEADER, Inches(0.8), Inches(1.35), width=Inches(11.7))

    add_card(slide3, Inches(0.8), Inches(2.2), Inches(11.7), Inches(4.7))
    tb3 = slide3.shapes.add_textbox(Inches(1.1), Inches(2.4), Inches(11.1), Inches(4.3))
    tf3 = tb3.text_frame
    tf3.word_wrap = True

    h_items = [
        ("• BỘ CHỌN VÙNG MIỀN (MB / MT / MN):", "KTV chọn đúng Miền phụ trách. Kho phiếu sẽ lọc và dùng chung độc lập theo từng Miền."),
        ("• 5 ĐÈN TRẠNG THÁI KẾT NỐI:", "Màu xanh (ONLINE): Hoạt động tốt. Màu đỏ (OFFLINE): Hết hạn cookie, KTV bấm vào để đăng nhập lại OneOSS/SAPC."),
        ("• NÚT 'LIVE LOG':", "Mở cửa sổ xem nhật ký tiến trình thời gian thực khi hệ thống quét phiếu hoặc đang tiền kiểm ngầm."),
        ("• QUẢN LÝ TÀI KHOẢN KTV:", "Góc trên bên phải hiển thị tên KTV đang đăng nhập, nút Đổi mật khẩu và Đăng xuất an toàn.")
    ]
    for idx, (lbl, val) in enumerate(h_items):
        p = tf3.paragraphs[0] if idx == 0 else tf3.add_paragraph()
        if idx > 0: p.space_before = Pt(12)
        r_l = p.add_run()
        r_l.text = f"{lbl} "
        r_l.font.name = "Arial"
        r_l.font.bold = True
        r_l.font.size = Pt(11)
        r_l.font.color.rgb = C_VNPT_BLUE
        r_v = p.add_run()
        r_v.text = val
        r_v.font.name = "Arial"
        r_v.font.size = Pt(10.5)
        r_v.font.color.rgb = C_DARK

    # =========================================================================
    # SLIDE 4: PHÂN HỆ 1 - MOBILE INTERNET (DATA)
    # =========================================================================
    slide4 = prs.slides.add_slide(blank_layout)
    add_header(slide4, "3. PHÂN HỆ 1: MOBILE INTERNET (3 THẺ CHI TIẾT)")

    if os.path.exists(IMG_EXPANDED):
        slide4.shapes.add_picture(IMG_EXPANDED, Inches(0.8), Inches(1.35), width=Inches(11.7))

    add_card(slide4, Inches(0.8), Inches(5.6), Inches(11.7), Inches(1.5))
    tb4 = slide4.shapes.add_textbox(Inches(1.0), Inches(5.7), Inches(11.3), Inches(1.3))
    tf4 = tb4.text_frame
    tf4.word_wrap = True

    p = tf4.paragraphs[0]
    p.text = "CẤU TRÚC 3 THẺ TÁC NGHIỆP MOBILE INTERNET (CLICK CHUỘT VÀO HÀNG PHIẾU ĐỂ MỞ):"
    p.font.name = "Arial"
    p.font.bold = True
    p.font.size = Pt(11)
    p.font.color.rgb = C_VNPT_BLUE

    p = tf4.add_paragraph()
    p.space_before = Pt(3)
    p.text = "• THẺ 1 (Trái): Tóm tắt AI (Gói khách báo, lỗi cụ thể, dung lượng báo còn, thiết bị, khu vực lỗi và địa chỉ Phường/Xã)\n• THẺ 2 (Giữa): Profile SAPC, Lịch sử phiên BTools và sóng trạm CEM (Đã bọc dòng wrapped text gọn gàng, không thanh cuộn ngang)\n• THẺ 3 (Phải): Ý kiến KTV, chọn mã nguyên nhân OneOSS và cụm nút thao tác (Chuyển 2.6 / Đóng 5.1 / Mở TTS)"
    p.font.name = "Arial"
    p.font.size = Pt(10)
    p.font.color.rgb = C_DARK

    # =========================================================================
    # SLIDE 5: QUY TẮC OCS: ODA_BIG_OCS TRONG PHÂN HỆ DATA
    # =========================================================================
    slide5 = prs.slides.add_slide(blank_layout)
    add_header(slide5, "4. ĐỐI SOÁT GÓI CƯỚC OCS: QUY TẮC BẮT BUỘC ODA_BIG_OCS")

    if os.path.exists(IMG_CARD2):
        slide5.shapes.add_picture(IMG_CARD2, Inches(0.8), Inches(1.35), width=Inches(5.7))

    add_card(slide5, Inches(6.8), Inches(1.35), Inches(5.7), Inches(5.6), bg_color=RGBColor(254, 252, 232), border_color=C_WARN)
    tb5 = slide5.shapes.add_textbox(Inches(7.0), Inches(1.5), Inches(5.3), Inches(5.2))
    tf5 = tb5.text_frame
    tf5.word_wrap = True

    p = tf5.paragraphs[0]
    p.text = "QUY TẮC BẮT LỖI GÓI OCS BẮT BUỘC:"
    p.font.name = "Arial"
    p.font.bold = True
    p.font.size = Pt(13)
    p.font.color.rgb = C_DANGER

    rules_ocs = [
        ("DANH MỤC GÓI OCS:", "D159SC, D159V, VD120N, VD120M, VD130, VD149, VD89, YOLO90, YOLO100M, YOLO125V, BIG60, BIG70, BIG80, D199G..."),
        ("DẤU HIỆU LỖI RỚT PROFILE:", "Khách báo có gói (VD: D159SC, còn 6GB) nhưng tra cứu SAPC CHỈ CÓ gói tiện ích ODA_MYTV mà THIẾU profile data ODA_BIG_OCS."),
        ("HƯỚNG XỬ LÝ CHUẨN CỦA KTV:", "'Chuyển IT/CCBS kiểm tra khai báo gói cước {tên gói} cho Khách hàng'."),
        ("CẤM TUYỆT ĐỐI:", "Không được tư vấn khách nạp tiền mua gói mới khi khách đã mua gói cước OCS!")
    ]
    for l, v in rules_ocs:
        p = tf5.add_paragraph()
        p.space_before = Pt(8)
        r_l = p.add_run()
        r_l.text = f"{l}\n"
        r_l.font.name = "Arial"
        r_l.font.bold = True
        r_l.font.size = Pt(10.5)
        r_l.font.color.rgb = C_VNPT_BLUE
        r_v = p.add_run()
        r_v.text = v
        r_v.font.name = "Arial"
        r_v.font.size = Pt(10)
        r_v.font.color.rgb = C_DARK

    # =========================================================================
    # SLIDE 6: PHÂN HỆ 2 - CUỘC GỌI & THOẠI (VOICE CALLS)
    # =========================================================================
    slide6 = prs.slides.add_slide(blank_layout)
    add_header(slide6, "5. PHÂN HỆ 2: CUỘC GỌI & THOẠI (VOICE CALLS)")

    if os.path.exists(IMG_VOICE):
        slide6.shapes.add_picture(IMG_VOICE, Inches(0.8), Inches(1.35), width=Inches(6.8))

    add_card(slide6, Inches(7.9), Inches(1.35), Inches(4.6), Inches(5.6))
    tb6 = slide6.shapes.add_textbox(Inches(8.1), Inches(1.5), Inches(4.2), Inches(5.2))
    tf6 = tb6.text_frame
    tf6.word_wrap = True

    p = tf6.paragraphs[0]
    p.text = "NGHIỆP VỤ TIỀN KIỂM THOẠI"
    p.font.name = "Arial"
    p.font.bold = True
    p.font.size = Pt(13)
    p.font.color.rgb = C_VNPT_BLUE

    items_voice = [
        ("• CỜ KHÓA CƯỚC THOẠI (NAM):", "NAM: 0 (Mở dịch vụ) - màu xanh lá.\nNAM: 1 (Khóa dịch vụ) - màu đỏ do nợ cước."),
        ("• HSS PROFILE THOẠI & VOLTE:", "Hiển thị profile HSS. Cảnh báo '⚠️ PROFILE LẠ' nếu chưa cấu hình VoLTE chuẩn."),
        ("• TỰ ĐỘNG ẨN BTOOLS LƯU LƯỢNG:", "Giao diện thoại tự động lược bỏ các thông tin data không liên quan, tập trung vào trạm phục vụ."),
        ("• ĐIỀU PHỐI BƯỚC 2.3 & 2.4:", "Hỗ trợ KTV duyệt nhanh chuyển bước 2.4 / 2.6 theo đúng luồng thoại OneOSS.")
    ]
    for l, v in items_voice:
        p = tf6.add_paragraph()
        p.space_before = Pt(8)
        r_l = p.add_run()
        r_l.text = f"{l}\n"
        r_l.font.name = "Arial"
        r_l.font.bold = True
        r_l.font.size = Pt(10)
        r_l.font.color.rgb = C_VNPT_ACCENT
        r_v = p.add_run()
        r_v.text = v
        r_v.font.name = "Arial"
        r_v.font.size = Pt(9.5)
        r_v.font.color.rgb = C_DARK

    # =========================================================================
    # SLIDE 7: PHÂN HỆ 3 - CHẶN SPAM NGOẠI MẠNG (NGHỊ ĐỊNH 91)
    # =========================================================================
    slide7 = prs.slides.add_slide(blank_layout)
    add_header(slide7, "6. PHÂN HỆ 3: CHẶN SPAM NGOẠI MẠNG (NGHỊ ĐỊNH 91)")

    if os.path.exists(IMG_SPAM):
        slide7.shapes.add_picture(IMG_SPAM, Inches(0.8), Inches(1.35), width=Inches(6.8))

    add_card(slide7, Inches(7.9), Inches(1.35), Inches(4.6), Inches(5.6))
    tb7 = slide7.shapes.add_textbox(Inches(8.1), Inches(1.5), Inches(4.2), Inches(5.2))
    tf7 = tb7.text_frame
    tf7.word_wrap = True

    p = tf7.paragraphs[0]
    p.text = "ĐỐI SOÁT CAM KẾT CCOS"
    p.font.name = "Arial"
    p.font.bold = True
    p.font.size = Pt(13)
    p.font.color.rgb = C_VNPT_BLUE

    items_spam = [
        ("• NHÀ MẠNG BỊ ẢNH HƯỞNG:", "Tự động phân tích và gắn huy hiệu nhà mạng chặn cuộc gọi (Viettel, MobiFone, Vietnamobile...)."),
        ("• ĐÃ CÓ BẢN CAM KẾT (XANH):", "Tìm thấy biên bản cam kết đính kèm trên CCOS. KTV chọn: 'Đã có bản cam kết, chuyển KTV kiểm tra mở chặn'."),
        ("• CHƯA CÓ BẢN CAM KẾT (ĐỎ):", "Chưa có cam kết trên CCOS. KTV phản hồi chuyển GDV yêu cầu khách ký cam kết không gọi rác."),
        ("• ĐIỀU PHỐI MỞ CHẶN LIÊN MẠNG:", "Rút ngắn thời gian xử lý khiếu nại chặn gọi rác theo đúng quy định.")
    ]
    for l, v in items_spam:
        p = tf7.add_paragraph()
        p.space_before = Pt(8)
        r_l = p.add_run()
        r_l.text = f"{l}\n"
        r_l.font.name = "Arial"
        r_l.font.bold = True
        r_l.font.size = Pt(10)
        r_l.font.color.rgb = C_VNPT_ACCENT
        r_v = p.add_run()
        r_v.text = v
        r_v.font.name = "Arial"
        r_v.font.size = Pt(9.5)
        r_v.font.color.rgb = C_DARK

    # =========================================================================
    # SLIDE 8: PHÂN HỆ 4 - TIN NHẮN (SMS) & CỔNG SMSC CDR
    # =========================================================================
    slide8 = prs.slides.add_slide(blank_layout)
    add_header(slide8, "7. PHÂN HỆ 4: TIN NHẮN (SMS) & ĐỐI SOÁT SMSC CDR")

    if os.path.exists(IMG_SMS):
        slide8.shapes.add_picture(IMG_SMS, Inches(0.8), Inches(1.35), width=Inches(6.8))

    add_card(slide8, Inches(7.9), Inches(1.35), Inches(4.6), Inches(5.6))
    tb8 = slide8.shapes.add_textbox(Inches(8.1), Inches(1.5), Inches(4.2), Inches(5.2))
    tf8 = tb8.text_frame
    tf8.word_wrap = True

    p = tf8.paragraphs[0]
    p.text = "TRA CỨU NHẬT KÝ SMSC"
    p.font.name = "Arial"
    p.font.bold = True
    p.font.size = Pt(13)
    p.font.color.rgb = C_VNPT_BLUE

    items_sms = [
        ("• KHIẾU NẠI SMS OTP & BRANDNAME:", "Khách không nhận được mã xác thực ngân hàng (Vietcombank, MB, Techcombank...), Zalo, Facebook."),
        ("• NÚT 'LẤY CDR' NHANH TRÊN HÀNG:", "Bấm 1 click chuột trực tiếp trên hàng phiếu để kéo kết quả nhật ký gửi tin tức thì."),
        ("• NÚT 'TRA CỨU SMSC CDR':", "Mở cửa sổ tra cứu toàn màn hình theo số điện thoại và khoảng thời gian tùy chọn."),
        ("• ĐỐI SOÁT TRỰC TIẾP GATEWAY:", "Kết nối thẳng trung tâm SMSC, xác định nguyên nhân trong vòng 30 giây thay vì 4 tiếng thủ công.")
    ]
    for l, v in items_sms:
        p = tf8.add_paragraph()
        p.space_before = Pt(8)
        r_l = p.add_run()
        r_l.text = f"{l}\n"
        r_l.font.name = "Arial"
        r_l.font.bold = True
        r_l.font.size = Pt(10)
        r_l.font.color.rgb = C_VNPT_ACCENT
        r_v = p.add_run()
        r_v.text = v
        r_v.font.name = "Arial"
        r_v.font.size = Pt(9.5)
        r_v.font.color.rgb = C_DARK

    # =========================================================================
    # SLIDE 9: CỬA SỔ TRA CỨU SMSC CDR CHI TIẾT
    # =========================================================================
    slide9 = prs.slides.add_slide(blank_layout)
    add_header(slide9, "8. CỬA SỔ TRA CỨU SMSC CDR CHI TIẾT (FULL SCREEN)")

    if os.path.exists(IMG_SMS_MODAL):
        slide9.shapes.add_picture(IMG_SMS_MODAL, Inches(0.8), Inches(1.35), width=Inches(6.8))

    add_card(slide9, Inches(7.9), Inches(1.35), Inches(4.6), Inches(5.6))
    tb9 = slide9.shapes.add_textbox(Inches(8.1), Inches(1.5), Inches(4.2), Inches(5.2))
    tf9 = tb9.text_frame
    tf9.word_wrap = True

    p = tf9.paragraphs[0]
    p.text = "CÁC CHỈ SỐ QUAN TRỌNG TRONG CDR"
    p.font.name = "Arial"
    p.font.bold = True
    p.font.size = Pt(13)
    p.font.color.rgb = C_VNPT_BLUE

    items_cdr = [
        ("• THỜI GIAN GỬI TIN:", "Giờ, phút, giây tin nhắn được đẩy vào hệ thống."),
        ("• ĐẦU SỐ DỊCH VỤ (SHORTCODE):", "Tên Brandname ngân hàng hoặc đầu số gửi tin."),
        ("• TRẠNG THÁI (DELIVERY STATUS):", "Delivered: Đã tới máy khách.\nExpired / Failed: Thất bại do tắt máy, sim đầy."),
        ("• MÃ LỖI KỸ THUẬT (ERROR CODE):", "Mã lỗi chuẩn viễn thông giúp KTV kết luận nguyên nhân chính xác 100%.")
    ]
    for l, v in items_cdr:
        p = tf9.add_paragraph()
        p.space_before = Pt(8)
        r_l = p.add_run()
        r_l.text = f"{l}\n"
        r_l.font.name = "Arial"
        r_l.font.bold = True
        r_l.font.size = Pt(10)
        r_l.font.color.rgb = C_VNPT_ACCENT
        r_v = p.add_run()
        r_v.text = v
        r_v.font.name = "Arial"
        r_v.font.size = Pt(9.5)
        r_v.font.color.rgb = C_DARK

    # =========================================================================
    # SLIDE 10: PHÂN HỆ 5 - CVQT (ROAMING) & SIM / ESIM
    # =========================================================================
    slide10 = prs.slides.add_slide(blank_layout)
    add_header(slide10, "9. PHÂN HỆ 5: CVQT (ROAMING) & SIM / ESIM")

    add_card(slide10, Inches(0.8), Inches(1.4), Inches(5.7), Inches(5.5))
    tb10_l = slide10.shapes.add_textbox(Inches(1.0), Inches(1.6), Inches(5.3), Inches(5.1))
    tf10_l = tb10_l.text_frame
    tf10_l.word_wrap = True

    p = tf10_l.paragraphs[0]
    p.text = "CHUYỂN VÙNG QUỐC TẾ (CVQT / ROAMING)"
    p.font.name = "Arial"
    p.font.bold = True
    p.font.size = Pt(13)
    p.font.color.rgb = C_VNPT_BLUE

    roam_pts = [
        ("• ĐỐI TƯỢNG:", "Khách VinaPhone ra nước ngoài phản ánh không có sóng hoặc không kết nối được mạng đối tác."),
        ("• NGHIỆP VỤ:", "Kiểm tra đăng ký dịch vụ CVQT (IR), gói cước Roaming (R15, R30...), trạng thái mở cước quốc tế."),
        ("• KẾT LUẬN:", "Khách chưa đăng ký dịch vụ, máy chưa bật Data Roaming, hoặc nhà mạng đối tác gián đoạn.")
    ]
    for l, v in roam_pts:
        p = tf10_l.add_paragraph()
        p.space_before = Pt(10)
        p.text = f"{l}\n{v}"
        p.font.name = "Arial"
        p.font.size = Pt(10)
        p.font.color.rgb = C_DARK

    add_card(slide10, Inches(6.8), Inches(1.4), Inches(5.7), Inches(5.5))
    tb10_r = slide10.shapes.add_textbox(Inches(7.0), Inches(1.6), Inches(5.3), Inches(5.1))
    tf10_r = tb10_r.text_frame
    tf10_r.word_wrap = True

    p = tf10_r.paragraphs[0]
    p.text = "DỊCH VỤ SIM / MULTISIM / ESIM"
    p.font.name = "Arial"
    p.font.bold = True
    p.font.size = Pt(13)
    p.font.color.rgb = C_VNPT_BLUE

    sim_pts = [
        ("• ĐỐI TƯỢNG:", "Khách lỗi khi nâng cấp SIM 4G, đổi phôi SIM, đăng ký MultiSIM hoặc quét mã QR eSIM."),
        ("• NGHIỆP VỤ:", "Kiểm tra số seri phôi SIM (IMSI, ICCID) trên HLR, trạng thái kích hoạt eSIM trên Core."),
        ("• KẾT LUẬN:", "Cấp lại mã QR eSIM hoặc hướng dẫn giao dịch viên đổi phôi SIM mới nếu phôi SIM vật lý bị lỗi.")
    ]
    for l, v in sim_pts:
        p = tf10_r.add_paragraph()
        p.space_before = Pt(10)
        p.text = f"{l}\n{v}"
        p.font.name = "Arial"
        p.font.size = Pt(10)
        p.font.color.rgb = C_DARK

    # =========================================================================
    # SLIDE 11: PHÂN HỆ 6 - QUẢN TRỊ KTV & PHÂN VÙNG SOC
    # =========================================================================
    slide11 = prs.slides.add_slide(blank_layout)
    add_header(slide11, "10. PHÂN HỆ 6: QUẢN TRỊ KTV & PHÂN VÙNG SOC")

    if os.path.exists(IMG_KTV_MGT):
        slide11.shapes.add_picture(IMG_KTV_MGT, Inches(0.8), Inches(1.35), width=Inches(6.8))

    add_card(slide11, Inches(7.9), Inches(1.35), Inches(4.6), Inches(5.6))
    tb11 = slide11.shapes.add_textbox(Inches(8.1), Inches(1.5), Inches(4.2), Inches(5.2))
    tf11 = tb11.text_frame
    tf11.word_wrap = True

    p = tf11.paragraphs[0]
    p.text = "TÍNH NĂNG QUẢN TRỊ SOC"
    p.font.name = "Arial"
    p.font.bold = True
    p.font.size = Pt(13)
    p.font.color.rgb = C_VNPT_BLUE

    items_mgt = [
        ("• PHÂN QUYỀN VÙNG MIỀN:", "Gán quyền KTV theo từng SOC (Miền Bắc, Miền Nam, Miền Trung)."),
        ("• QUẢN LÝ PHIÊN ONEOSS:", "Cập nhật tài khoản, mật khẩu và token truy cập cho các phân vùng."),
        ("• CẤU HÌNH AI & CHẨN ĐOÁN:", "Tùy chỉnh kịch bản tự động, chu kỳ quét nền 60s."),
        ("• THEO DÕI HIỆU NĂNG:", "Giám sát số lượng phiếu xử lý và thời gian đáp ứng hệ thống.")
    ]
    for l, v in items_mgt:
        p = tf11.add_paragraph()
        p.space_before = Pt(8)
        r_l = p.add_run()
        r_l.text = f"{l}\n"
        r_l.font.name = "Arial"
        r_l.font.bold = True
        r_l.font.size = Pt(10)
        r_l.font.color.rgb = C_VNPT_ACCENT
        r_v = p.add_run()
        r_v.text = v
        r_v.font.name = "Arial"
        r_v.font.size = Pt(9.5)
        r_v.font.color.rgb = C_DARK

    # =========================================================================
    # SLIDE 12: DANH MỤC 73 MÃ NGUYÊN NHÂN ONEOSS
    # =========================================================================
    slide12 = prs.slides.add_slide(blank_layout)
    add_header(slide12, "11. DANH MỤC 73 MÃ NGUYÊN NHÂN SỰ CỐ ONEOSS")

    if os.path.exists(IMG_DROPDOWN):
        slide12.shapes.add_picture(IMG_DROPDOWN, Inches(0.8), Inches(1.35), width=Inches(5.7))

    add_card(slide12, Inches(6.8), Inches(1.35), Inches(5.7), Inches(5.6))
    tb12 = slide12.shapes.add_textbox(Inches(7.0), Inches(1.5), Inches(5.3), Inches(5.2))
    tf12 = tb12.text_frame
    tf12.word_wrap = True

    p = tf12.paragraphs[0]
    p.text = "CÁCH CHỌN MÃ NGUYÊN NHÂN"
    p.font.name = "Arial"
    p.font.bold = True
    p.font.size = Pt(13)
    p.font.color.rgb = C_VNPT_BLUE

    items_cause = [
        ("1. HỆ THỐNG GỢI Ý SẴN:", "Dựa trên tiền kiểm, bot đã tự điền sẵn mã lỗi OneOSS phù hợp nhất."),
        ("2. TÌM KIẾM THEO TỪ KHÓA:", "Gõ từ khóa vào ô (VD: 'gói cước', 'thiết bị', 'bình thường'...), danh sách 73 mã lỗi sẽ lọc ngay."),
        ("3. MÃ SỐ RÕ RÀNG:", "Mỗi nguyên nhân có mã ID OneOSS đi kèm (ví dụ: #2012, #2015...)."),
        ("4. TỰ ĐỘNG LƯU:", "Khi KTV chọn nguyên nhân, hệ thống tự động lưu vào phiếu.")
    ]
    for l, v in items_cause:
        p = tf12.add_paragraph()
        p.space_before = Pt(8)
        r_l = p.add_run()
        r_l.text = f"{l}\n"
        r_l.font.name = "Arial"
        r_l.font.bold = True
        r_l.font.size = Pt(10)
        r_l.font.color.rgb = C_VNPT_ACCENT
        r_v = p.add_run()
        r_v.text = v
        r_v.font.name = "Arial"
        r_v.font.size = Pt(9.5)
        r_v.font.color.rgb = C_DARK

    # =========================================================================
    # SLIDE 13: CỤM NÚT THAO TÁC - CHUYỂN 2.6 VS ĐÓNG 5.1
    # =========================================================================
    slide13 = prs.slides.add_slide(blank_layout)
    add_header(slide13, "12. CỤM NÚT THAO TÁC: CHUYỂN 2.6 VS ĐÓNG 5.1")

    if os.path.exists(IMG_BTNS):
        slide13.shapes.add_picture(IMG_BTNS, Inches(0.8), Inches(1.35), width=Inches(11.7))

    btns_comp = [
        ("NÚT 'CHUYỂN 2.6'", "Màu xanh nhạt viền xanh", C_VNPT_ACCENT,
         "KHI NÀO DÙNG:\nPhiếu đang ở bước 2.4 và KTV muốn chuyển sang bước kỹ thuật 2.6 nhưng CHƯA ĐÓNG HẲN.\n\nTÁC DỤNG THỰC TẾ:\nCập nhật kết luận và nâng bước phiếu lên 2.6 trên OneOSS. Phiếu vẫn mở để theo dõi."),
        ("NÚT 'ĐÓNG 5.1'", "Màu xanh đậm chữ trắng", C_VNPT_BLUE,
         "KHI NÀO DÙNG:\nĐã xác định rõ nguyên nhân (sóng bình thường, hết data, VPN, lỗi gói cước...) và muốn ĐÓNG HOÀN TẤT.\n\nTÁC DỤNG THỰC TẾ:\nGửi mã nguyên nhân, giải trình kỹ thuật và ĐÓNG DỨT ĐIỂM phiếu trên OneOSS."),
        ("NÚT 'MỞ TTS'", "Màu xám chữ xanh", C_MUTED,
         "KHI NÀO DÙNG:\nMuốn đối soát sâu trên giao diện OneOSS gốc hoặc xử lý ca khiếu nại phức tạp.\n\nTÁC DỤNG THỰC TẾ:\nMở tab mới đưa KTV thẳng vào trang phiếu OneOSS.")
    ]
    for i, (b_name, b_sub, color, b_desc) in enumerate(btns_comp):
        left = Inches(0.8 + i * 4.0)
        add_card(slide13, left, Inches(2.2), Inches(3.7), Inches(4.7))

        tb = slide13.shapes.add_textbox(left + Inches(0.2), Inches(2.4), Inches(3.3), Inches(4.3))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = b_name
        p.font.name = "Arial"
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = color

        p = tf.add_paragraph()
        p.space_before = Pt(3)
        p.text = b_sub
        p.font.name = "Arial"
        p.font.size = Pt(9.5)
        p.font.italic = True
        p.font.color.rgb = C_MUTED

        p = tf.add_paragraph()
        p.space_before = Pt(12)
        p.text = b_desc
        p.font.name = "Arial"
        p.font.size = Pt(10)
        p.font.color.rgb = C_DARK

    # =========================================================================
    # SLIDE 14: QUY CHUẨN BẮT BUỘC ĐÓNG 5.1 - CẬP NHẬT PHƯỜNG / XÃ
    # =========================================================================
    slide14 = prs.slides.add_slide(blank_layout)
    add_header(slide14, "13. QUY CHUẨN BẮT BUỘC ĐÓNG 5.1: CẬP NHẬT PHƯỜNG / XÃ")

    if os.path.exists(IMG_WARD):
        slide14.shapes.add_picture(IMG_WARD, Inches(0.8), Inches(1.35), width=Inches(5.7))

    add_card(slide14, Inches(6.8), Inches(1.35), Inches(5.7), Inches(5.6), bg_color=RGBColor(254, 252, 232), border_color=C_WARN)
    tb14 = slide14.shapes.add_textbox(Inches(7.0), Inches(1.5), Inches(5.3), Inches(5.2))
    tf14 = tb14.text_frame
    tf14.word_wrap = True

    p = tf14.paragraphs[0]
    p.text = "QUY CHUẨN ĐỊA BÀN 2 CẤP BẮT BUỘC"
    p.font.name = "Arial"
    p.font.bold = True
    p.font.size = Pt(13)
    p.font.color.rgb = C_WARN

    ward_rules = [
        ("QUY CHUẨN DUY NHẤT 2 CẤP:", "Định dạng: [Phường/Xã/Thị trấn], [Tỉnh/Thành phố]\nTUYỆT ĐỐI KHÔNG lấy cấp thôn, ấp, bản, khóm, tổ dân phố (như Ấp 1, Thôn 3, Tổ 5...)."),
        ("👉 CÁCH 1 (1 CLICK TỰ ĐỘNG - KHUYÊN DÙNG):", "Nhìn dòng 'Radio:' ở Thẻ 2. Nếu trạm BTS có tọa độ GPS, bấm NÚT MŨI TÊN XANH (↑) bên cạnh. Hệ thống tự động điền địa bàn chuẩn từ Google Maps (Port 1708)!"),
        ("👉 CÁCH 2 (SỬA TRỰC TIẾP TRÊN Ô TEXT):", "Nếu khách báo dạng 'ấp lương phú C, xã Hòa Lạc, Tỉnh Đồng Tháp' -> KTV xóa bỏ phần ấp, giữ lại đúng: 'Xã Hòa Lạc, Tỉnh Đồng Tháp'."),
        ("DANH MỤC 34 TỈNH/TP MỚI:", "Bắt buộc theo bảng danh mục mới (MB: 18 Tỉnh/TP, MN: 9 Tỉnh/TP, MT: 7 Tỉnh/TP).")
    ]
    for l, v in ward_rules:
        p = tf14.add_paragraph()
        p.space_before = Pt(8)
        r_l = p.add_run()
        r_l.text = f"{l}\n"
        r_l.font.name = "Arial"
        r_l.font.bold = True
        r_l.font.size = Pt(10)
        r_l.font.color.rgb = C_DANGER if "QUY CHUẨN" in l else C_VNPT_BLUE
        r_v = p.add_run()
        r_v.text = v
        r_v.font.name = "Arial"
        r_v.font.size = Pt(9.5)
        r_v.font.color.rgb = C_DARK

    # =========================================================================
    # SLIDE 15: DANH MỤC 34 TỈNH / THÀNH PHỐ MỚI VÀ 3 MIỀN
    # =========================================================================
    slide15 = prs.slides.add_slide(blank_layout)
    add_header(slide15, "14. BẢNG 34 TỈNH / THÀNH PHỐ MỚI & PHÂN BỔ 3 MIỀN")

    regions_34 = [
        ("MIỀN BẮC (MB)", "18 ĐƠN VỊ HÀNH CHÍNH", C_VNPT_BLUE,
         "• 2 Thành phố: TP. Hà Nội, TP. Hải Phòng (gồm Hải Dương).\n• 16 Tỉnh: Cao Bằng, Lạng Sơn, Lai Châu, Điện Biên, Sơn La, Quảng Ninh, Tuyên Quang (gồm Hà Giang), Lào Cai (gồm Yên Bái), Thái Nguyên (gồm Bắc Kạn), Phú Thọ (gồm Vĩnh Phúc, Hòa Bình), Bắc Ninh (gồm Bắc Giang), Hưng Yên (gồm Thái Bình), Ninh Bình (gồm Hà Nam, Nam Định), Thanh Hóa, Nghệ An, Hà Tĩnh."),
        ("MIỀN NAM (MN)", "9 ĐƠN VỊ HÀNH CHÍNH", C_SUCCESS,
         "• 2 Thành phố: TP. Hồ Chí Minh (gồm Bình Dương, Bà Rịa - Vũng Tàu), TP. Cần Thơ (gồm Sóc Trăng, Hậu Giang).\n• 7 Tỉnh: Lâm Đồng (gồm Đắk Nông, Bình Thuận), Đồng Nai (gồm Bình Phước), Tây Ninh (gồm Long An), Vĩnh Long (gồm Bến Tre, Trà Vinh), Đồng Tháp (gồm Tiền Giang), Cà Mau (gồm Bạc Liêu), An Giang (gồm Kiên Giang & Đặc khu Phú Quốc)."),
        ("MIỀN TRUNG (MT)", "7 ĐƠN VỊ HÀNH CHÍNH", C_WARN,
         "• 2 Thành phố: TP. Huế, TP. Đà Nẵng (gồm Quảng Nam).\n• 5 Tỉnh: Quảng Trị (gồm Quảng Bình), Quảng Ngãi (gồm Kon Tum), Gia Lai (gồm Bình Định), Khánh Hòa (gồm Ninh Thuận), Đắk Lắk (gồm Phú Yên).")
    ]
    for i, (r_name, r_sub, color, r_list) in enumerate(regions_34):
        left = Inches(0.8 + i * 4.0)
        add_card(slide15, left, Inches(1.4), Inches(3.7), Inches(5.5))

        tb = slide15.shapes.add_textbox(left + Inches(0.2), Inches(1.6), Inches(3.3), Inches(5.1))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = r_name
        p.font.name = "Arial"
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = color

        p = tf.add_paragraph()
        p.space_before = Pt(3)
        p.text = r_sub
        p.font.name = "Arial"
        p.font.size = Pt(9.5)
        p.font.bold = True
        p.font.color.rgb = C_MUTED

        p = tf.add_paragraph()
        p.space_before = Pt(10)
        p.text = r_list
        p.font.name = "Arial"
        p.font.size = Pt(9.5)
        p.font.color.rgb = C_DARK

    # =========================================================================
    # SLIDE 16: QUY TRÌNH 5 BƯỚC TÁC NGHIỆP CHO NEWBIE
    # =========================================================================
    slide16 = prs.slides.add_slide(blank_layout)
    add_header(slide16, "15. QUY TRÌNH 5 BƯỚC TÁC NGHIỆP TỪ A ĐẾN Z CHO NEWBIE")

    steps_5 = [
        ("BƯỚC 1", "ĐĂNG NHẬP & CHỌN MIỀN", "Truy cập http://10.155.59.158:1234, chọn đúng Miền (MB/MT/MN). Kiểm tra các đèn trạng thái đều màu xanh lá."),
        ("BƯỚC 2", "CHỌN PHÂN HỆ & LÀM MỚI", "Chọn phân hệ cần làm (Data, Thoại, Spam, SMS). Bấm tab 'Cần xử lý', bấm nút 'Làm mới' tải phiếu OneOSS."),
        ("BƯỚC 3", "MỞ CHI TIẾT & TIỀN KIỂM", "Click chuột vào hàng phiếu để mở 3 Thẻ. Bấm '⚡ Tiền kiểm' (hoặc 'Lấy CDR' với tin nhắn) để nạp dữ liệu mạng."),
        ("BƯỚC 4", "KIỂM TRA Ý KIẾN & NGUYÊN NHÂN", "Đọc kết luận Thẻ 3. Kiểm tra ô combobox Nguyên nhân sự cố đã được bot gợi ý đúng mã lỗi OneOSS chưa."),
        ("BƯỚC 5", "CHUẨN HÓA ĐỊA BÀN & ĐÓNG PHIẾU", "Kiểm tra địa bàn đúng [Phường/Xã], [Tỉnh/TP] (bấm nút mũi tên xanh ↑). Bấm nút 'Đóng 5.1' hoàn tất phiếu!")
    ]
    for i, (b_num, b_title, b_desc) in enumerate(steps_5):
        top_y = Inches(1.35 + i * 1.15)
        add_card(slide16, Inches(0.8), top_y, Inches(11.7), Inches(1.0))

        badge = slide16.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), top_y + Inches(0.15), Inches(1.4), Inches(0.7))
        badge.fill.solid()
        badge.fill.fore_color.rgb = C_VNPT_BLUE
        badge.line.color.rgb = C_VNPT_BLUE
        tf_b = badge.text_frame
        tf_b.margin_left = tf_b.margin_top = tf_b.margin_right = tf_b.margin_bottom = 0
        p_b = tf_b.paragraphs[0]
        p_b.alignment = PP_ALIGN.CENTER
        p_b.text = b_num
        p_b.font.name = "Arial"
        p_b.font.size = Pt(13)
        p_b.font.bold = True
        p_b.font.color.rgb = C_WHITE

        tb = slide16.shapes.add_textbox(Inches(2.6), top_y + Inches(0.1), Inches(9.7), Inches(0.8))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p = tf.paragraphs[0]
        p.text = b_title
        p.font.name = "Arial"
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = C_VNPT_ACCENT

        p = tf.add_paragraph()
        p.space_before = Pt(2)
        p.text = b_desc
        p.font.name = "Arial"
        p.font.size = Pt(10)
        p.font.color.rgb = C_DARK

    # =========================================================================
    # SLIDE 17: CÁC LỖI THƯỜNG GẶP VÀ CÁCH XỬ LÝ (FAQ)
    # =========================================================================
    slide17 = prs.slides.add_slide(blank_layout)
    add_header(slide17, "16. CÁC TÌNH HUỐNG THƯỜNG GẶP VÀ CÁCH XỬ LÝ (FAQ)")

    faqs = [
        ("1. BẤM 'LÀM MỚI' NHƯNG KHÔNG CÓ PHIẾU?", 
         "Kiểm tra đèn trạng thái TTS Mới ở Header. Nếu màu đỏ nghĩa là hết hạn phiên cookie OneOSS, bấm vào để đăng nhập lại."),
        ("2. KHÔNG THẤY PHIẾU CỦA ĐỒNG NGHIỆP CÙNG MIỀN?", 
         "Đảm bảo ô chọn Vùng Miền ở góc trên bên trái đang chọn đúng Miền (MB/MN/MT) và kết nối mạng nội bộ ổn định."),
        ("3. GÓI OCS (D159SC...) BÁO THIẾU DATA TRÊN SAPC?", 
         "Gói OCS bắt buộc phải có profile ODA_BIG_OCS. Nếu chỉ có ODA_MYTV, chọn nguyên nhân 'Lỗi do gói cước' và hướng xử lý 'Chuyển IT/CCBS kiểm tra khai báo'."),
        ("4. ĐÓNG 5.1 BỊ BÁO LỖI ĐỊA BÀN PHƯỜNG / XÃ?", 
         "Kiểm tra lại ô địa chỉ, chỉ giữ lại đúng 2 cấp [Phường/Xã], [Tỉnh/TP]. Tuyệt đối xóa bỏ cấp thôn, ấp, bản, tổ dân phố.")
    ]
    for i, (q, a) in enumerate(faqs):
        row = i // 2
        col = i % 2
        left = Inches(0.8 + col * 6.0)
        top = Inches(1.4 + row * 2.8)
        add_card(slide17, left, top, Inches(5.7), Inches(2.5))

        tb = slide17.shapes.add_textbox(left + Inches(0.2), top + Inches(0.2), Inches(5.3), Inches(2.1))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = q
        p.font.name = "Arial"
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = C_DANGER

        p = tf.add_paragraph()
        p.space_before = Pt(6)
        p.text = a
        p.font.name = "Arial"
        p.font.size = Pt(10)
        p.font.color.rgb = C_DARK

    # =========================================================================
    # SLIDE 18: LỜI CẢM ƠN & THÔNG TIN HỖ TRỢ
    # =========================================================================
    slide18 = prs.slides.add_slide(blank_layout)
    bg18 = slide18.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg18.fill.solid()
    bg18.fill.fore_color.rgb = C_VNPT_BLUE
    bg18.line.color.rgb = C_VNPT_BLUE

    card_end = slide18.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(2.0), Inches(1.8), Inches(9.333), Inches(3.9))
    card_end.fill.solid()
    card_end.fill.fore_color.rgb = C_WHITE
    card_end.line.color.rgb = C_WHITE

    tb18 = slide18.shapes.add_textbox(Inches(2.3), Inches(2.2), Inches(8.7), Inches(3.1))
    tf18 = tb18.text_frame
    tf18.word_wrap = True

    p = tf18.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = "CHÚC CÁC ANH EM KỸ THUẬT VIÊN\nTÁC NGHIỆP HIỆU QUẢ VÀ CHÍNH XÁC!"
    p.font.name = "Arial"
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = C_VNPT_BLUE

    p = tf18.add_paragraph()
    p.alignment = PP_ALIGN.CENTER
    p.space_before = Pt(15)
    p.text = "TẬP ĐOÀN BƯU CHÍNH VIỄN THÔNG VIỆT NAM (VNPT)\nKHỐI KỸ THUẬT VIỄN THÔNG - TRUNG TÂM ĐIỀU HÀNH MẠNG"
    p.font.name = "Arial"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = C_VNPT_ACCENT

    p = tf18.add_paragraph()
    p.alignment = PP_ALIGN.CENTER
    p.space_before = Pt(20)
    p.text = "Hệ thống VNPT TTS Precheck  |  Cổng 1234  |  Hỗ trợ KTV 24/7"
    p.font.name = "Arial"
    p.font.size = Pt(10)
    p.font.italic = True
    p.font.color.rgb = C_MUTED

    output_pptx = r"c:\Users\lequa\OneDrive - VNPT\python\PAKH_PRECHECK\SLIDE_HUONG_DAN_SU_DUNG_VNPT_TTS_PRECHECK.pptx"
    prs.save(output_pptx)
    print("SUCCESS: File PowerPoint TOAN DIEN 18 SLIDE da duoc tao thanh cong tai:", output_pptx)

if __name__ == "__main__":
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass
    create_full_slide_deck()
