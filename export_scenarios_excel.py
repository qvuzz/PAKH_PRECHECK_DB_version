# export_scenarios_excel.py
# Xuat danh sach toan bo kich ban dong tu dong Mobile Internet:
# Bao gom: Dong 2.6, Dong 5.1 va Can KTV Dong thu cong

import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUTPUT_FILE = "DANH_MUC_KICH_BAN_DONG_PHIEU_MOBILE_INTERNET.xlsx"

wb = Workbook()
wb.remove(wb.active)  # Xoa sheet mac dinh

# Dinh nghia mau sac & kieu dang chuan Doanh nghiep VNPT
COLOR_PRIMARY = "004B8D"       # Xanh VNPT dac trung
COLOR_WHITE = "FFFFFF"
COLOR_HEADER_BG = "004B8D"
COLOR_26_BG = "E2EFDA"         # Xanh la nhat cho Dong 2.6
COLOR_26_BORDER = "375623"
COLOR_51_BG = "D9E1F2"         # Xanh duong nhat cho Dong 5.1
COLOR_51_BORDER = "1F4E79"
COLOR_MANUAL_BG = "FCE4D6"     # Cam nhat cho Thu cong
COLOR_MANUAL_BORDER = "C65911"
COLOR_GRAY_LIGHT = "F2F2F2"

font_title = Font(name="Segoe UI", size=14, bold=True, color=COLOR_WHITE)
font_subtitle = Font(name="Segoe UI", size=10, italic=True, color="D9D9D9")
font_header = Font(name="Segoe UI", size=10, bold=True, color=COLOR_WHITE)
font_group_header = Font(name="Segoe UI", size=11, bold=True)
font_body = Font(name="Segoe UI", size=9)
font_body_bold = Font(name="Segoe UI", size=9, bold=True)
font_code = Font(name="Consolas", size=9, bold=True)

thin_border_side = Side(border_style="thin", color="D9D9D9")
border_cell = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
border_header = Border(
    left=Side(border_style="thin", color="003366"),
    right=Side(border_style="thin", color="003366"),
    top=Side(border_style="thin", color="003366"),
    bottom=Side(border_style="medium", color="002244")
)

align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
align_top_left = Alignment(horizontal="left", vertical="top", wrap_text=True)
align_top_center = Alignment(horizontal="center", vertical="top", wrap_text=True)

fill_header = PatternFill(start_color=COLOR_HEADER_BG, end_color=COLOR_HEADER_BG, fill_type="solid")
fill_zebra = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")

def apply_table_styles(ws, max_cols, header_row=3):
    ws.views.sheetView[0].showGridLines = True
    for col in range(1, max_cols + 1):
        cell = ws.cell(row=header_row, column=col)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = border_header
    ws.row_dimensions[header_row].height = 28

def auto_adjust_columns(ws, max_cols, min_w=12, max_w=50):
    for col in range(1, max_cols + 1):
        max_len = 0
        col_letter = get_column_letter(col)
        for row in range(3, ws.max_row + 1):
            val = ws.cell(row=row, column=col).value
            if val:
                lines = str(val).split("\n")
                line_max = max(len(l) for l in lines)
                if line_max > max_len:
                    max_len = line_max
        w = max(min_w, min(max_len + 3, max_w))
        ws.column_dimensions[col_letter].width = w

# ==============================================================================
# SHEET 1: TONG QUAN & QUY TRINH
# ==============================================================================
ws_overview = wb.create_sheet(title="TỔNG QUAN HỆ THỐNG")

ws_overview.merge_cells("A1:G1")
cell_t = ws_overview.cell(row=1, column=1, value="HỆ THỐNG ĐÓNG PHIẾU TỰ ĐỘNG MOBILE INTERNET - VNPT VTT / ONEOSS")
cell_t.font = font_title
cell_t.fill = fill_header
cell_t.alignment = align_center
ws_overview.row_dimensions[1].height = 40

ws_overview.merge_cells("A2:G2")
cell_sub = ws_overview.cell(row=2, column=1, value="Quy chuẩn nghiệp vụ phân loại kịch bản: Hướng Đóng 2.6, Hướng Đóng 5.1 và Rào chắn bảo vệ KTV Đóng thủ công")
cell_sub.font = font_subtitle
cell_sub.fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
cell_sub.alignment = align_center
ws_overview.row_dimensions[2].height = 24

overview_stats = [
    ("Nhóm Kịch bản", "Bước trên OneOSS TTS", "Tỷ lệ xử lý thực tế", "Cơ chế xử lý", "Mục tiêu & Phạm vi áp dụng"),
    ("HƯỚNG ĐÓNG 2.6 (ĐÓNG DỨT ĐIỂM)", "Bước 2.6 - Đóng phiếu Trên TTS", "85% - 90% lượng phiếu", "Tự động 100% qua REST API", "Hoàn tất đóng phiếu dứt điểm ngay trên hệ thống MSC/418. Dành cho các lỗi do gói cước, bóp băng thông, thiết bị khách hàng, SIM, profile Core hoặc mạng lưới bình thường."),
    ("HƯỚNG ĐÓNG 5.1 (CHUYỂN ĐỊA BÀN VTT)", "Bước 5.1 - Xây dựng PA xử lý", "7% - 10% lượng phiếu", "Chuyển bước bán tự động / Tự động định tuyến", "Chuyển VNPT Tỉnh / Đội kỹ thuật viễn thông địa bàn để đi đo kiểm hiện trường, tối ưu vùng phủ sóng, xử lý trạm suy hao hoặc nghẽn tải."),
    ("CẦN KTV ĐÓNG THỦ CÔNG", "Giữ nguyên bước 2.3 / 2.4", "3% - 5% lượng phiếu", "KTV kiểm tra & chốt thủ công", "Rào chắn an toàn: Các phiếu mở lại nhiều lần (reopen), phản ánh lỗi ứng dụng riêng lẻ (Zalo, TikTok), sai lệch địa bàn Phường/Xã hoặc nghi vấn tính cước.")
]

ws_overview.cell(row=4, column=1, value="1. BẢNG THỐNG KÊ PHÂN PHỐI LUỒNG XỬ LÝ PHIẾU").font = Font(name="Segoe UI", size=11, bold=True, color="004B8D")
for r_idx, row_data in enumerate(overview_stats, start=5):
    for c_idx, val in enumerate(row_data, start=1):
        cell = ws_overview.cell(row=r_idx, column=c_idx, value=val)
        cell.border = border_cell
        if r_idx == 5:
            cell.font = font_header
            cell.fill = PatternFill(start_color="203764", end_color="203764", fill_type="solid")
            cell.alignment = align_center
        else:
            cell.font = font_body
            cell.alignment = align_top_left if c_idx == 5 else align_top_center
            if c_idx == 1:
                cell.font = font_body_bold
                if "2.6" in val:
                    cell.fill = PatternFill(start_color=COLOR_26_BG, end_color=COLOR_26_BG, fill_type="solid")
                elif "5.1" in val:
                    cell.fill = PatternFill(start_color=COLOR_51_BG, end_color=COLOR_51_BG, fill_type="solid")
                else:
                    cell.fill = PatternFill(start_color=COLOR_MANUAL_BG, end_color=COLOR_MANUAL_BG, fill_type="solid")
    ws_overview.row_dimensions[r_idx].height = 42 if r_idx > 5 else 26

# Luong xu ly
rules_text = [
    ("2. NGUYÊN TẮC VẬN HÀNH & ĐIỀU KIỆN TIÊN QUYẾT BẢO ĐẢM AN TOÀN TOÀN TRÌNH:", ""),
    ("Nguyên tắc 1: Luồng chuyển bước chuẩn OneOSS", "Phiếu xuất hiện lần đầu ở bước 2.4 (hoặc 2.3). Nếu đủ điều kiện đóng 2.6, hệ thống chuyển sang bước 2.6 (chốt mạng lưới). Khi phiếu tái xuất hiện ở 2.6, hệ thống kích hoạt API close-ticket đóng dứt điểm."),
    ("Nguyên tắc 2: Ràng buộc địa bàn nghiêm ngặt khi đóng 5.1", "Hệ thống kiểm tra bắt buộc phiếu phải có Phường/Xã và Tỉnh/TP chuẩn (chuẩn hóa 2 cấp), đối chiếu khớp với trạm CEM và kết quả Port 1708. Tự động cập nhật Lĩnh vực = 'Chất lượng mạng' (ID: 71) trước khi chuyển bước 5.1. Nếu thiếu địa bàn -> Chặn chuyển 5.1 để tránh trôi phiếu sai đơn vị."),
    ("Nguyên tắc 3: Rào chắn bảo vệ Phiếu khiếu nại nhiều lần", "Bất kỳ phiếu nào có trường 'Số lần mở lại' (reopen_count) > 0 hoặc nội dung có ghi nhận mở lại -> Tuyệt đối không đóng tự động, đưa vào danh sách Cần KTV đóng thủ công."),
    ("Nguyên tắc 4: Rào chắn lỗi ứng dụng cụ thể", "Khách hàng phản ánh lỗi ứng dụng riêng lẻ (Zalo, Facebook, TikTok, VNeID, app ngân hàng...) trong khi các phiên data chung vẫn phát sinh -> Tuyệt đối không đóng tự động, giao KTV hỗ trợ chuyên biệt.")
]

start_r = 11
for item in rules_text:
    title, desc = item
    if not desc:
        ws_overview.cell(row=start_r, column=1, value=title).font = Font(name="Segoe UI", size=11, bold=True, color="004B8D")
        start_r += 1
    else:
        ws_overview.cell(row=start_r, column=1, value=title).font = font_body_bold
        ws_overview.cell(row=start_r, column=1).alignment = align_top_left
        ws_overview.cell(row=start_r, column=1).border = border_cell
        ws_overview.cell(row=start_r, column=1).fill = PatternFill(start_color="EDF2F8", end_color="EDF2F8", fill_type="solid")
        
        ws_overview.merge_cells(start_row=start_r, start_column=2, end_row=start_r, end_column=5)
        c_desc = ws_overview.cell(row=start_r, column=2, value=desc)
        c_desc.font = font_body
        c_desc.alignment = align_top_left
        for col in range(2, 6):
            ws_overview.cell(row=start_r, column=col).border = border_cell
        ws_overview.row_dimensions[start_r].height = 36
        start_r += 1

ws_overview.column_dimensions["A"].width = 32
ws_overview.column_dimensions["B"].width = 28
ws_overview.column_dimensions["C"].width = 22
ws_overview.column_dimensions["D"].width = 28
ws_overview.column_dimensions["E"].width = 50

# ==============================================================================
# SHEET 2: KICH BAN DONG 2.6 (DONG DUT DIEM)
# ==============================================================================
ws_26 = wb.create_sheet(title="ĐÓNG 2.6 (ĐÓNG DỨT ĐIỂM)")

ws_26.merge_cells("A1:I1")
c = ws_26.cell(row=1, column=1, value="DANH MỤC KỊCH BẢN ĐÓNG 2.6 - HOÀN TẤT ĐÓNG PHIẾU DỨT ĐIỂM TRÊN TTS")
c.font = font_title
c.fill = PatternFill(start_color="1E4620", end_color="1E4620", fill_type="solid")
c.alignment = align_center
ws_26.row_dimensions[1].height = 36

ws_26.merge_cells("A2:I2")
c = ws_26.cell(row=2, column=1, value="Áp dụng cho các sự cố đã xác minh nguyên nhân Core HSS / SAPC / Thiết bị / Gói cước hoặc Mạng lưới bình thường - Tự động đóng trực tiếp qua REST API (Tỷ lệ 85% - 90%)")
c.font = font_subtitle
c.fill = PatternFill(start_color="2D6A4F", end_color="2D6A4F", fill_type="solid")
c.alignment = align_center
ws_26.row_dimensions[2].height = 22

headers_26 = [
    "STT", "Mã KB", "Tên Kịch Bản Chi Tiết", "Điều Kiện Nhận Diện Hạ Tầng & Dữ Liệu",
    "Nguyên Nhân Sự Cố (OneOSS Dropdown)", "Nội Dung Phân Tích Kỹ Thuật (KTV Comment)",
    "Hướng Xử Lý Khuyên Dùng (Action Plan)", "Mã Lỗi / Căn Cứ", "Mức Độ Tự Động"
]
for idx, h in enumerate(headers_26, 1):
    cell = ws_26.cell(row=3, column=idx, value=h)
apply_table_styles(ws_26, len(headers_26), 3)

data_26 = [
    # 1. Bop bang thong
    (
        "KB26_01", "Thuê bao bị bóp băng thông tốc độ cao (Hạ băng thông)",
        "BTools ghi nhận các mã dịch vụ bóp băng thông chính thức VNPT: 10002 (64/64), 10003 (128/64), 10005 (512/256), 10010 (1024/1024), 10011 (2048/2048), 10012 (3072/3072), 10013 (2048/1024), 10014 (128/128) hoặc dải 100xx.",
        "Thông tin đầu vào chưa chính xác, trùng lặp",
        "Hệ thống BTools ghi nhận mã dịch vụ giới hạn băng thông {MÃ_CODE} (tốc độ {TỐC_ĐỘ} Kbps). Thuê bao đã sử dụng hết dung lượng tốc độ cao trong chu kỳ gói cước và bị hạ băng thông theo đúng chính sách gói.",
        "Thông báo khách hàng đã sử dụng hết dung lượng tốc độ cao của gói cước. Tư vấn khách hàng mua thêm dung lượng data tốc độ cao để tiếp tục truy cập Internet.",
        "KC_03_THROTTLED", "Tự động 100%"
    ),
    # 2. Mang luoi dam bao (Phien lon sau tiep nhan)
    (
        "KB26_02", "Mạng lưới đảm bảo - Thuê bao sử dụng tốt sau tiếp nhận",
        "Sau mốc thời gian khách hàng phản ánh sự cố (Incident Time), BTools ghi nhận phiên dữ liệu phát sinh bình thường > 10MB (hoặc tổng lưu lượng lớn), throughput đạt chuẩn 4G/5G, không rớt mạng.",
        "Mạng lưới đảm bảo, khách hàng sử dụng dịch vụ bình thường",
        "Kiểm tra sau thời điểm phản ánh ({GIỜ_PHẢN_ÁNH}), thuê bao vẫn duy trì kết nối mạng di động ổn định, phát sinh các phiên truy cập dữ liệu lớn (phiên đạt {DUNG_LƯỢNG_MAX}MB), tốc độ truyền tải đảm bảo chất lượng kỹ thuật.",
        "Mạng lưới hoạt động bình thường. Đóng phiếu dứt điểm, không cần can thiệp hạ tầng.",
        "LEVEL_1_STATUS", "Tự động 100%"
    ),
    # 3. Goi het han
    (
        "KB26_03", "Gói cước data đã hết hạn sử dụng",
        "Đối soát hệ thống SAPC: Toàn bộ các gói cước Data của thuê bao đều đã hết hạn sử dụng (Expire Date < Thời điểm phản ánh). Thuê bao không còn gói data chu kỳ.",
        "Lỗi do gói cước",
        "Kiểm tra hệ thống quản lý chính sách cước SAPC, các gói data của thuê bao đã hết hạn sử dụng trước thời điểm phản ánh. Thuê bao hiện không còn gói cước dữ liệu hợp lệ.",
        "Tư vấn khách hàng gia hạn hoặc đăng ký gói cước data mới để tiếp tục truy cập mạng.",
        "SAPC_EXPIRED", "Tự động 100%"
    ),
    # 4. Chua dang ky goi
    (
        "KB26_04", "Thuê bao chưa đăng ký gói cước Data",
        "Hồ sơ SAPC không ghi nhận bất kỳ gói cước Data thương mại nào, tài khoản chỉ có gói mặc định M0 hoặc tài khoản chính không đủ điều kiện kết nối data.",
        "Lỗi do gói cước",
        "Thuê bao chưa đăng ký gói cước Mobile Internet, hệ thống ghi nhận dịch vụ dữ liệu chưa được kích hoạt gói thương mại.",
        "Hướng dẫn khách hàng cú pháp đăng ký gói cước data VinaPhone phù hợp với nhu cầu.",
        "NO_PACKAGE", "Tự động 100%"
    ),
    # 5. Chi co goi PAYGO
    (
        "KB26_05", "Thuê bao chỉ có gói cước tính cước theo dung lượng (PAYGO)",
        "SAPC ghi nhận thuê bao chỉ kích hoạt gói PAYGO (Pay As You Go), không có gói lưu lượng data trọn gói hoặc gói ngày/tháng.",
        "Lỗi do gói cước",
        "Thuê bao chỉ cài đặt gói PAYGO tính cước theo dung lượng thực tế, dễ phát sinh chặn kết nối khi hết tiền tài khoản chính.",
        "Tư vấn khách hàng đăng ký các gói cước data ưu đãi trọn gói (VD, BIG, YOLO...) để tối ưu chi phí.",
        "PAYGO_ONLY", "Tự động 100%"
    ),
    # 6. Chi co goi tien ich - thieu data internet
    (
        "KB26_06", "Chỉ có gói tiện ích ứng dụng (Add-on) - Thiếu Data Internet",
        "SAPC chỉ ghi nhận các gói data tiện ích theo ứng dụng (như DIP_TikTok, DIP_Youtube, MyTV...) nhưng không có gói data Internet thông thường (không có data truy cập web/ứng dụng khác).",
        "Lỗi do gói cước",
        "Thuê bao chỉ đăng ký gói cước tiện ích ứng dụng chuyên biệt, không có lưu lượng data truy cập Internet thông thường. Khi truy cập các ứng dụng ngoài gói sẽ không kết nối được.",
        "Giải thích cho khách hàng về phạm vi sử dụng của gói tiện ích và tư vấn đăng ký thêm gói data Internet nền.",
        "ADDON_NO_DATA", "Tự động 100%"
    ),
    # 7. Goi VD2 thieu PAYGO
    (
        "KB26_07", "Lỗi gói cước ngày VD2 - Thiếu cấu hình PAYGO",
        "SAPC ghi nhận thuê bao đăng ký gói ngày VD2 nhưng thiếu gói nền PAYGO dẫn tới OCS không định tuyến mở phiên data.",
        "Lỗi do gói cước",
        "Thuê bao đăng ký gói cước VD2 nhưng hệ thống thiếu chính sách gói nền PAYGO đồng bộ.",
        "Hỗ trợ đồng bộ lại chính sách cước PAYGO trên hệ thống SAPC cho thuê bao.",
        "PKG_VD2_ERR", "Tự động 100%"
    ),
    # 8. Goi Tha ga thieu PAYGO
    (
        "KB26_08", "Lỗi gói cước Thả Ga - Thiếu cấu hình PAYGO",
        "SAPC ghi nhận thuê bao có gói Thả Ga nhưng thiếu profile PAYGO đi kèm.",
        "Lỗi do gói cước",
        "Thuê bao có gói Thả Ga nhưng thiếu chính sách nền PAYGO trên hệ thống tính cước.",
        "Cập nhật lại gói cước nền PAYGO trên SAPC.",
        "PKG_THAGA_ERR", "Tự động 100%"
    ),
    # 9. Tru cuoc ngoai goi PAYGO
    (
        "KB26_09", "Phát sinh cước ngoài gói PAYGO (Mã 3001)",
        "BTools ghi nhận mã dịch vụ 3001 phát sinh cước ngoài gói khi gói data chính hết dung lượng.",
        "Lỗi do gói cước",
        "Thuê bao phát sinh phiên dữ liệu theo hình thức trừ cước ngoài gói (mã 3001) sau khi đã sử dụng hết dung lượng gói cước chính.",
        "Giải thích cho khách hàng về việc cước ngoài gói phát sinh do gói chính hết dung lượng; tư vấn mua thêm gói phụ.",
        "PAYGO_3001", "Tự động 100%"
    ),
    # 10. VPN / 1.1.1.1
    (
        "KB26_10", "Khách hàng sử dụng VPN / 1.1.1.1 / Cloudflare WARP",
        "CEM App Usage phát hiện ứng dụng VPN (1.1.1.1, WARP, OpenVPN, ExpressVPN...). BTools các phiên downlink duy trì cố định bất thường trong dải 5MB - 11.5MB (chặn trần VPN) trên RAT_TYPE 1 hoặc 6.",
        "Do thiết bị đầu cuối",
        "Hệ thống phát hiện thiết bị khách hàng đang bật ứng dụng mạng riêng ảo (VPN / 1.1.1.1 / Cloudflare WARP). Ứng dụng VPN làm chuyển hướng và suy hao băng thông quốc tế, gây hiện tượng nghẽn mạng cục bộ trên máy.",
        "Liên hệ hỗ trợ khách hàng tắt ứng dụng VPN/1.1.1.1 đã cài đặt trên máy và kiểm tra lại kết nối trực tiếp qua mạng VinaPhone.",
        "KC_06_VPN_OR_DEVICE_ISSUE", "Tự động 100%"
    ),
    # 11. Thiet bi di nhieu noi bi loi
    (
        "KB26_11", "Lỗi thiết bị đầu cuối - Đi nhiều nơi đều phản ánh lỗi",
        "Nội dung phản ánh khách báo 'đi đâu cũng bị', 'nhiều nơi không vào được'; CEM ghi nhận thuê bao kết nối qua nhiều trạm BTS tại các quận/tỉnh khác nhau nhưng đều suy hao đồng đều -> Lỗi phần cứng thiết bị.",
        "Do thiết bị đầu cuối",
        "Khách hàng phản ánh lỗi xảy ra ở nhiều vị trí địa lý khác nhau. Dữ liệu trạm CEM ghi nhận thuê bao kết nối qua nhiều trạm BTS phân tán, không có trạm nào phát sinh sự cố diện rộng. Xác định nguyên nhân do thiết bị hoặc SIM.",
        "Hướng dẫn khách hàng khởi động lại máy, cài đặt lại APN hoặc mang máy/SIM qua điểm giao dịch kiểm tra thiết bị.",
        "DEVICE_MULTI_LOC", "Tự động 100%"
    ),
    # 12. Tat du lieu di dong / Treo thiet bi
    (
        "KB26_12", "Thiết bị tắt Dữ liệu di động hoặc treo Data",
        "Hệ thống HSS Core và SAPC bình thường, nhưng BTools hoàn toàn không ghi nhận bất kỳ phiên data nào phát sinh trong 5 ngày gần nhất (không có cả phiên hệ thống).",
        "Do thiết bị đầu cuối",
        "Hệ thống mạng lưới và gói cước hoàn toàn bình thường, nhưng thiết bị không phát sinh bất kỳ yêu cầu kết nối dữ liệu nào lên mạng lưới. Nghi ngờ thiết bị tắt Dữ liệu di động hoặc bị lỗi treo phân hệ mạng.",
        "Hướng dẫn khách hàng bật tính năng 'Dữ liệu di động' (Cellular Data) trên máy, khởi động lại thiết bị để tái lập kết nối.",
        "NO_DATA_SESSIONS", "Tự động 100%"
    ),
    # 13. Khong co luu luong dang ke (Toan ma quan ly phien)
    (
        "KB26_13", "Không có lưu lượng đáng kể - Chỉ phát sinh mã quản trị phiên",
        "BTools chỉ ghi nhận các mã quản trị hệ thống (300, 302, 330, 2042), không có mã data thương mại; tổng dung lượng phát sinh < 1MB trong suốt N ngày.",
        "Do thiết bị đầu cuối",
        "Hệ thống chỉ ghi nhận các bản tin quản trị phiên định kỳ (mã 300, 302, 2042), không phát sinh lưu lượng data thực tế do người dùng sử dụng. Thiết bị chưa kích hoạt kết nối ứng dụng mạng.",
        "Hướng dẫn khách hàng kiểm tra lại các ứng dụng chạy ngầm, bật/tắt lại chế độ máy bay để refresh phiên kết nối.",
        "KC_04_PACKAGE_OR_DEVICE_HANG", "Tự động 100%"
    ),
    # 14. MS PURGED (Tat may nhieu ngay)
    (
        "KB26_14", "Thuê bao tắt máy dài ngày (HSS Sub State = MS PURGED)",
        "Kiểm tra HSS Core: Trạng thái thuê bao ghi nhận 'MS PURGED', thuê bao không duy trì đăng ký vị trí với mạng lưới trong nhiều ngày liên tục.",
        "Do thiết bị đầu cuối",
        "Hệ thống HSS Core ghi nhận trạng thái thuê bao MS PURGED (thiết bị tắt máy hoặc không hoạt động trên mạng lưới trong thời gian dài).",
        "Khách hàng vui lòng bật nguồn thiết bị và gắn SIM để mạng lưới tự động kích hoạt lại trạng thái phục vụ.",
        "HSS_MS_PURGED", "Tự động 100%"
    ),
    # 15. Chua khai bao Profile 4G
    (
        "KB26_15", "Chưa khai báo Profile 4G trên HSS Core",
        "HSS Core Profile rỗng hoặc chưa kích hoạt tham số LTE/4G (HSS Profile = 0/null), thiết bị chỉ kết nối được qua sóng 2G/3G.",
        "Lỗi Profile thuê bao",
        "Kiểm tra hệ thống HSS Core, thuê bao chưa được khai báo profile dịch vụ 4G LTE, dẫn tới thiết bị không thể kết nối mạng 4G.",
        "Chuyển bộ phận IT / Khai thác mạng kiểm tra, kích hoạt bổ sung profile 4G cho thuê bao trên HSS.",
        "KC_07_NO_4G_PROFILE", "Tự động 100%"
    ),
    # 16. Profile la
    (
        "KB26_16", "Thuê bao có Profile lạ trên hệ thống HSS",
        "HSS Profile chứa các tham số dịch vụ đặc thù hoặc profile từ 3 chữ số trở lên không thuộc danh mục gói tiêu chuẩn VinaPhone.",
        "Lỗi Profile thuê bao",
        "Thuê bao đang được gán profile dịch vụ đặc biệt trên hệ thống HSS, có thể gây xung đột quyền truy cập dữ liệu di động thông thường.",
        "Chuyển IT kiểm tra và chuẩn hóa lại profile dịch vụ cho thuê bao theo đúng gói cước đăng ký.",
        "STRANGE_PROFILE", "Tự động 100%"
    ),
    # 17. Khoa chan GPRS (NAM = 1)
    (
        "KB26_17", "Thuê bao bị khóa chặn dịch vụ dữ liệu GPRS (NAM = 1)",
        "HSS Core ghi nhận cờ NAM (Network Access Mode) = 1 (bị chặn dịch vụ dữ liệu GPRS/Packet Data do nợ cước hoặc khóa 1 chiều).",
        "Lỗi do VNPT-VinaPhone khai báo dịch vụ cho khách hàng",
        "Hệ thống HSS Core ghi nhận cờ khóa truy cập mạng dữ liệu (NAM = 1). Dịch vụ GPRS của thuê bao đang bị khóa tạm thời.",
        "Kiểm tra tình trạng thanh toán cước/nạp tiền của thuê bao và mở lại cờ NAM = 0 trên hệ thống.",
        "NAM_BLOCKED_1", "Tự động 100%"
    ),
    # 18. HSS chua co 5G
    (
        "KB26_18", "Khách hàng phản ánh 5G nhưng Profile Core chưa mở 5G",
        "Khách hàng phản ánh không bắt được 5G, nhưng trên Core HSS thuê bao chưa được mở profile 5G NSA/SA.",
        "Lỗi do VNPT-VinaPhone khai báo dịch vụ cho khách hàng",
        "Khách hàng phản ánh dịch vụ 5G, tuy nhiên hệ thống Core HSS chưa kích hoạt cấu hình dịch vụ 5G cho thuê bao.",
        "Bổ sung cấu hình mở dịch vụ 5G cho thuê bao trên hệ thống Core.",
        "HSS_NO_5G", "Tự động 100%"
    ),
    # 19. Khach hang theo doi them
    (
        "KB26_19", "Khách hàng theo dõi thêm (Sự cố tạm thời / Không có vị trí cụ thể)",
        "Khách phản ánh chập chờn nhưng không cung cấp vị trí địa chỉ cụ thể; trước tiếp nhận có data, sau tiếp nhận chưa đủ phiên kiểm chứng; hoặc phản ánh trong thời gian ngắn đã ổn định.",
        "Khách hàng theo dõi thêm",
        "Khách hàng phản ánh sự cố tạm thời hoặc không xác định vị trí lỗi cố định. Mạng lưới tại khu vực tiếp nhận ghi nhận thông số tải bình thường. Có thể do nghẽn mạng cục bộ tại thời điểm khách hàng di chuyển.",
        "Nhờ khách hàng tiếp tục theo dõi thêm dịch vụ. Nếu sự cố tái diễn tại vị trí cố định, vui lòng liên hệ tổng đài kèm địa chỉ chi tiết.",
        "FOLLOW_UP", "Tự động 100%"
    ),
    # 20. Luu luong yeu khong tap trung 1 cell
    (
        "KB26_20", "Lưu lượng yếu phân tán (Không khoanh vùng 1 khu vực)",
        "Lưu lượng data trong khoảng 1MB - 10MB nhưng phân tán đều qua nhiều trạm CEM khác nhau, không có trạm nào vượt 50% lưu lượng.",
        "Khách hàng theo dõi thêm",
        "Lưu lượng dữ liệu của khách hàng phân tán qua nhiều trạm phát sóng khác nhau trong ngày, không ghi nhận suy hao cục bộ tại trạm cụ thể.",
        "Nhờ khách hàng theo dõi thêm chất lượng dịch vụ khi di chuyển giữa các khu vực.",
        "WEAK_DATA_DISPERSED", "Tự động 100%"
    ),
    # 21. Loi goi cuoc - Sai Service ID
    (
        "KB26_21", "Lỗi chính sách gói cước - Sai lệch Service ID",
        "SAPC ghi nhận gói cước hợp lệ nhưng BTools ghi nhận mã dịch vụ (Service ID) không khớp với mã cước chuẩn của gói.",
        "Lỗi do gói cước",
        "Hệ thống phát hiện có sự sai lệch giữa mã chính sách cước trên SAPC và mã dịch vụ phiên dữ liệu thực tế trên BTools (Service ID Mismatch).",
        "Chuyển IT đồng bộ lại chính sách gói cước trên OCS/SAPC cho thuê bao.",
        "SERVICE_ID_MISMATCH", "Tự động 100%"
    ),
    # 22. Goi Home / Nghen bang thong chia se
    (
        "KB26_22", "Gói cước tích hợp HOME / Nghẽn băng thông chia sẻ",
        "Thuê bao sử dụng gói cước gia đình/tích hợp HOME chia sẻ dung lượng với các thành viên khác trong nhóm.",
        "Lỗi do gói cước",
        "Thuê bao sử dụng gói cước tích hợp HOME chia sẻ dung lượng dữ liệu. Dung lượng chung của nhóm có thể đã bị các thuê bao thành viên khác sử dụng hết.",
        "Kiểm tra dung lượng chia sẻ của nhóm HOME và thông báo chủ nhóm mua thêm dung lượng nếu cần.",
        "PKG_HOME_SHARING", "Tự động 100%"
    )
]

for row_idx, r_data in enumerate(data_26, start=4):
    stt = row_idx - 3
    row_vals = [stt] + list(r_data)
    for col_idx, val in enumerate(row_vals, start=1):
        cell = ws_26.cell(row=row_idx, column=col_idx, value=val)
        cell.border = border_cell
        cell.font = font_body
        if col_idx in (1, 2, 9):
            cell.alignment = align_top_center
            if col_idx == 2:
                cell.font = font_code
            elif col_idx == 9:
                cell.font = font_body_bold
                cell.fill = PatternFill(start_color=COLOR_26_BG, end_color=COLOR_26_BG, fill_type="solid")
        else:
            cell.alignment = align_top_left
            if col_idx == 3:
                cell.font = font_body_bold
    ws_26.row_dimensions[row_idx].height = 48
    if row_idx % 2 == 1:
        for col_idx in range(1, len(row_vals) + 1):
            if col_idx != 9:
                ws_26.cell(row=row_idx, column=col_idx).fill = fill_zebra

auto_adjust_columns(ws_26, len(headers_26))

# ==============================================================================
# SHEET 3: KICH BAN DONG 5.1 (CHUYEN DIA BAN VTT)
# ==============================================================================
ws_51 = wb.create_sheet(title="ĐÓNG 5.1 (CHUYỂN ĐỊA BÀN VTT)")

ws_51.merge_cells("A1:I1")
c = ws_51.cell(row=1, column=1, value="DANH MỤC KỊCH BẢN ĐÓNG 5.1 - CHUYỂN BƯỚC XÂY DỰNG PA XỬ LÝ (KỸ THUẬT ĐỊA BÀN VTT)")
c.font = font_title
c.fill = PatternFill(start_color="0F2027", end_color="0F2027", fill_type="solid")
c.alignment = align_center
ws_51.row_dimensions[1].height = 36

ws_51.merge_cells("A2:I2")
c = ws_51.cell(row=2, column=1, value="Áp dụng cho các sự cố suy hao vô tuyến, mép vùng phủ sóng hoặc trạm BTS cảnh báo lỗi - Bắt buộc kiểm tra ràng buộc Phường/Xã chuẩn và Lĩnh vực 71 (CLM)")
c.font = font_subtitle
c.fill = PatternFill(start_color="203A43", end_color="203A43", fill_type="solid")
c.alignment = align_center
ws_51.row_dimensions[2].height = 22

headers_51 = [
    "STT", "Mã KB", "Tên Kịch Bản Chi Tiết", "Dấu Hiệu Nhận Diện Kỹ Thuật (Hạ Tầng / CEM / BTools)",
    "Ràng Buộc Địa Bàn Bắt Buộc", "Lĩnh Vực TTS", "Nội Dung Phân Tích Kỹ Thuật Chuyển Địa Bàn",
    "Hướng Xử Lý Khuyên Dùng (Chuyển VTT)", "Đơn Vị Phối Hợp Xử Lý"
]
for idx, h in enumerate(headers_51, 1):
    cell = ws_51.cell(row=3, column=idx, value=h)
apply_table_styles(ws_51, len(headers_51), 3)

data_51 = [
    (
        "KB51_01", "Bắt sóng 4G kém / Chập chờn (Thiết bị rớt về 2G/3G chiếm đa số)",
        "BTools ghi nhận có phiên 4G (RAT 6) nhưng số phiên rớt về 2G/3G (RAT 1, 2) chiếm đa số (hoặc tổng phiên >= 3 và tỷ lệ rớt 2G/3G >= 70%). Biểu hiện thuê bao ở mép vùng phủ sóng 4G.",
        "Bắt buộc đúng Phường/Xã và Tỉnh/TP chuẩn mới (2 cấp). Đối chiếu khớp với địa chỉ trạm CEM phục vụ.",
        "Chất lượng mạng (ID: 71)",
        "Thiết bị có kết nối 4G nhưng tần suất rớt về sóng 2G/3G chiếm đa số trong các ngày gần đây (tỷ lệ phiên 2G/3G đạt {TỶ_LỆ}%). Biểu hiện thuê bao đang hoạt động tại khu vực sóng 4G yếu hoặc mép vùng phủ sóng trạm BTS.",
        "Chuyển kỹ thuật địa bàn VTT kiểm tra chất lượng vùng phủ sóng 4G tại khu vực khách hàng phản ánh ({ĐỊA_BÀN_CHUẨN}).",
        "VTT Tỉnh/TP / Tổ Kỹ thuật địa bàn"
    ),
    (
        "KB51_02", "Không bắt được sóng 4G (Dù HSS Profile Core đã kích hoạt 4G)",
        "HSS Core Profile đã mở 4G (HSS Profile khác rỗng/0), nhưng BTools trong suốt N ngày gần nhất hoàn toàn KHÔNG có phiên 4G (RAT 6/7 = 0%), toàn bộ kết nối chỉ bắt 2G/3G (RAT 0, 1, 2).",
        "Bắt buộc đúng Phường/Xã và Tỉnh/TP chuẩn mới. Trùng khớp với khu vực trạm phục vụ.",
        "Chất lượng mạng (ID: 71)",
        "Hệ thống HSS Core đã khai báo profile dịch vụ 4G hoàn chỉnh, nhưng thiết bị hoàn toàn không bắt được sóng 4G trong các ngày gần đây xuyên suốt. Khu vực phản ánh có thể bị lõm sóng 4G.",
        "Chuyển kỹ thuật địa bàn VTT kiểm tra mức phát sóng trạm 4G tại khu vực thuê bao phản ánh; hỗ trợ hướng dẫn KH kiểm tra thiết bị.",
        "VTT Tỉnh/TP / Tổ Kỹ thuật địa bàn"
    ),
    (
        "KB51_03", "Sóng 4G chập chờn / Yếu tại một địa chỉ cụ thể",
        "Khách hàng phản ánh sóng chập chờn khoanh vùng tại 1 địa chỉ cố định (nhà riêng, cơ quan). CEM ghi nhận mức tín hiệu RSRP/RSRQ thấp kéo dài.",
        "Bắt buộc chuẩn hóa Phường/Xã theo địa chỉ cố định khách hàng phản ánh.",
        "Chất lượng mạng (ID: 71)",
        "Khách hàng phản ánh sóng 4G chập chờn khoanh vùng tại địa chỉ cố định ({ĐỊA_CHỈ}). Dữ liệu CEM ghi nhận chỉ số chất lượng phủ sóng tại khu vực trạm chưa đạt mức tối ưu.",
        "Chuyển kỹ thuật địa bàn VTT đo kiểm mức sóng trong nhà và ngoài trời tại vị trí khách hàng phản ánh; đề xuất giải pháp phát sóng tối ưu.",
        "VTT Tỉnh/TP / Tổ Kỹ thuật địa bàn"
    ),
    (
        "KB51_04", "Lưu lượng yếu kéo dài tập trung tại một Cell phát sóng (Dominant Cell > 50%)",
        "CEM ghi nhận 1 trạm/cell phục vụ ưu thế chiếm > 50% tổng lưu lượng của thuê bao; lưu lượng data trong ngày yếu và tốc độ thấp kéo dài tại riêng Cell này.",
        "Bắt buộc đúng Phường/Xã vị trí đặt trạm BTS theo dữ liệu Port 1708.",
        "Chất lượng mạng (ID: 71)",
        "Thuê bao hoạt động chủ yếu dưới vùng phủ sóng của Cell {TÊN_CELL} (chiếm {TỶ_LỆ}% lưu lượng), tốc độ truyền tải ghi nhận mức thấp kéo dài.",
        "Chuyển kỹ thuật địa bàn VTT kiểm tra tải và thông số góc ngẩng anten / công suất phát sóng của trạm {TÊN_CELL}.",
        "VTT Tỉnh/TP / Đội Vận hành Khai thác trạm"
    ),
    (
        "KB51_05", "Trạm BTS có cảnh báo sự cố kỹ thuật / Bảo dưỡng (FMS Radio Alarm)",
        "Đối soát hệ thống FMS và Port 1708 ghi nhận trạm BTS phục vụ thuê bao đang có cảnh báo sự cố cảnh báo vô tuyến (Radio Loss, VSWR cao, mất liên lạc tạm thời) hoặc lịch bảo dưỡng.",
        "Lấy theo Phường/Xã của trạm BTS đang bị sự cố trên FMS.",
        "Chất lượng mạng (ID: 71)",
        "Trạm BTS {MÃ_TRẠM} tại khu vực phục vụ thuê bao ghi nhận có sự cố vô tuyến / bảo dưỡng thiết bị tại thời điểm phát sinh phản ánh.",
        "Chuyển kỹ thuật địa bàn VTT / Trung tâm Điều hành mạng theo dõi xử lý dứt điểm cảnh báo trạm {MÃ_TRẠM} và khôi phục dịch vụ.",
        "Trung tâm Mạng lưới / Đội Vận hành Trạm"
    ),
    (
        "KB51_06", "Nghẽn tải vô tuyến cục bộ giờ cao điểm (Cell Congestion)",
        "Throughput giờ cao điểm suy giảm mạnh, trong khi giờ thấp điểm tốc độ đạt bình thường; trạm CEM ghi nhận số lượng thuê bao đồng thời cao.",
        "Bắt buộc chuẩn hóa Phường/Xã vị trí đặt trạm theo Port 1708.",
        "Chất lượng mạng (ID: 71)",
        "Khu vực trạm BTS phục vụ khách hàng ghi nhận hiện tượng nghẽn tải lưu lượng vô tuyến cục bộ vào các khung giờ cao điểm.",
        "Chuyển kỹ thuật địa bàn VTT kiểm tra mở rộng băng thông phát sóng / cân bằng tải sang các cell lân cận.",
        "VTT Tỉnh/TP / Phòng Kỹ thuật Hạ tầng"
    )
]

for row_idx, r_data in enumerate(data_51, start=4):
    stt = row_idx - 3
    row_vals = [stt] + list(r_data)
    for col_idx, val in enumerate(row_vals, start=1):
        cell = ws_51.cell(row=row_idx, column=col_idx, value=val)
        cell.border = border_cell
        cell.font = font_body
        if col_idx in (1, 2, 6, 9):
            cell.alignment = align_top_center
            if col_idx == 2:
                cell.font = font_code
            elif col_idx == 6:
                cell.font = font_body_bold
                cell.fill = PatternFill(start_color=COLOR_51_BG, end_color=COLOR_51_BG, fill_type="solid")
        else:
            cell.alignment = align_top_left
            if col_idx == 3:
                cell.font = font_body_bold
    ws_51.row_dimensions[row_idx].height = 54
    if row_idx % 2 == 1:
        for col_idx in range(1, len(row_vals) + 1):
            if col_idx != 6:
                ws_51.cell(row=row_idx, column=col_idx).fill = fill_zebra

auto_adjust_columns(ws_51, len(headers_51))

# ==============================================================================
# SHEET 4: CAN KTV DONG THU CONG
# ==============================================================================
ws_manual = wb.create_sheet(title="CẦN KTV ĐÓNG THỦ CÔNG")

ws_manual.merge_cells("A1:H1")
c = ws_manual.cell(row=1, column=1, value="DANH MỤC TRƯỜNG HỢP CẦN KTV ĐÓNG THỦ CÔNG - RÀO CHẮN AN TOÀN NGHIỆP VỤ")
c.font = font_title
c.fill = PatternFill(start_color="8A2BE2", end_color="8A2BE2", fill_type="solid")
c.alignment = align_center
ws_manual.row_dimensions[1].height = 36

ws_manual.merge_cells("A2:H2")
c = ws_manual.cell(row=2, column=1, value="Bao gồm các trường hợp rủi ro cao: Phiếu mở lại nhiều lần, phản ánh lỗi ứng dụng riêng lẻ, sai lệch địa bàn hoặc nghi ngờ lỗi tính cước OCS - Tuyệt đối không tự động đóng")
c.font = font_subtitle
c.fill = PatternFill(start_color="5E2B97", end_color="5E2B97", fill_type="solid")
c.alignment = align_center
ws_manual.row_dimensions[2].height = 22

headers_manual = [
    "STT", "Mã Tình Huống", "Tên Tình Huống Cần Can Thiệp Thủ Công",
    "Dấu Hiệu Nhận Diện Trên Hệ Thống", "Lý Do Chặn Tự Động Đóng (Nguyên Tắc An Toàn)",
    "Rủi Ro Nếu Tự Động Đóng", "Hành Động Bắt Buộc Của Kỹ Thuật Viên", "Khuyến Nghị Hướng Xử Lý"
]
for idx, h in enumerate(headers_manual, 1):
    cell = ws_manual.cell(row=3, column=idx, value=h)
apply_table_styles(ws_manual, len(headers_manual), 3)

data_manual = [
    (
        "MAN_01", "Phiếu đã mở lại nhiều lần (reopen_count > 0)",
        "Trường 'Số lần mở lại' trên TTS Mới >= 1 (thông tin mở lại TTS); khách hàng khiếu nại lặp lại nhiều lần do chưa hài lòng với kết quả xử lý trước đó.",
        "Khách hàng đang bức xúc gay gắt, đã khiếu nại nhiều lần. Tự động đóng sẽ vi phạm quy chế chất lượng CSKH của VNPT.",
        "Khách hàng khiếu nại vượt cấp lên Sở TT&TT, Cục Viễn thông hoặc gọi tổng đài gay gắt.",
        "KTV bắt buộc phải gọi điện trực tiếp cho khách hàng, xác minh chi tiết hiện tượng tại nhà và phối hợp đo kiểm thực tế.",
        "Xử lý dứt điểm nguyên nhân gốc rễ, có biên bản làm việc hoặc ghi âm thỏa thuận trước khi đóng thủ công."
    ),
    (
        "MAN_02", "Phản ánh lỗi ứng dụng cụ thể (Zalo, TikTok, Facebook, YouTube, App Ngân hàng...)",
        "Khách hàng báo mạng vẫn có nhưng không vào được 1 ứng dụng cụ thể; BTools vẫn ghi nhận data chung phát sinh đều đặn; AI phân tích nhận diện 'lỗi ứng dụng cụ thể'.",
        "Dữ liệu BTools chỉ đo lưu lượng tổng qua APN, không phân tích được lỗi máy chủ riêng của ứng dụng bên thứ ba (server Zalo/Meta lỗi, DNS app, khóa tài khoản...).",
        "Đóng nhầm 'Mạng lưới đảm bảo' trong khi khách vẫn không vào được app, gây bức xúc.",
        "KTV liên hệ hỏi khách phiên bản ứng dụng, hướng dẫn xóa cache ứng dụng, kiểm tra cấp quyền mạng di động trong cài đặt máy hoặc đổi DNS.",
        "Hướng dẫn khách hàng kiểm tra cập nhật app trên App Store / Google Play; nếu server app bảo trì thì giải thích cho khách."
    ),
    (
        "MAN_03", "Gói cước còn hạn nhưng hoàn toàn không dùng được từ khi đăng ký",
        "SAPC ghi nhận gói cước vừa đăng ký thành công, còn nguyên quota dung lượng, nhưng từ thời điểm đăng ký tới nay BTools = 0MB (không phát sinh data).",
        "Nghi ngờ hệ thống tính cước OCS bị treo Profile hoặc lỗi đồng bộ gói giữa OCS và SAPC/PGW.",
        "Thuê bao bị mất tiền cước nhưng không sử dụng được dịch vụ, nguy cơ khiếu nại cước.",
        "KTV kiểm tra log tính cước trên OCS, kiểm tra cờ NAM, kích hoạt reset profile cước trên SAPC.",
        "Reset chính sách gói trên SAPC hoặc chuyển phòng Tính cước OCS bù lưu lượng cho khách."
    ),
    (
        "MAN_04", "Địa bàn Phường/Xã trên TTS Mới bị bỏ trống hoặc sai khác với Trạm CEM thực tế",
        "Phiếu thuộc hướng chất lượng sóng (5.1) nhưng trên TTS Mới chưa chọn Tỉnh/TP hoặc Phường/Xã; hoặc địa bàn trên phiếu là Tỉnh A nhưng trạm CEM thực tế phát sóng tại Tỉnh B.",
        "Nếu tự động chuyển bước 5.1 sẽ chuyển nhầm đơn vị viễn thông tỉnh khác, gây tình trạng đùn đẩy phiếu giữa các đơn vị và trễ SLA.",
        "Phiếu bị trả lại hoặc trôi nổi sai địa bàn, vi phạm chỉ số SLA đóng phiếu viễn thông.",
        "KTV kiểm tra vị trí trạm CEM thực tế qua Port 1708, cập nhật lại đúng Phường/Xã và Tỉnh/TP trên TTS Mới trước khi chuyển 5.1.",
        "Chuẩn hóa địa bàn 2 cấp chuẩn theo Port 1708, sau đó mới bấm chuyển 5.1."
    ),
    (
        "MAN_05", "Nghi ngờ SIM bị lỗi vật lý hoặc SIM cũ chưa đổi sang USIM 4G",
        "HSS Core có 4G, trạm phát 4G tốt nhưng máy chỉ nhận 2G; khách báo lắp máy khác vẫn không nhận sóng 4G; thời gian sử dụng SIM > 5 năm.",
        "Hệ thống không thể thay thế việc kiểm tra vi mạch vật lý của thẻ SIM ngoài quầy giao dịch.",
        "Đóng phiếu không giải quyết được lỗi phần cứng SIM của khách hàng.",
        "KTV hướng dẫn khách hàng mang căn cước công dân ra điểm giao dịch VinaPhone gần nhất để kiểm tra đổi thẻ SIM 4G miễn phí.",
        "Ghi chú chuyển giao dịch viên hỗ trợ đổi SIM vật lý."
    ),
    (
        "MAN_06", "Dịch vụ ngoài Mobile Internet (Thoại VoLTE/CSFB, SMS Brandname, MNP...)",
        "Phiếu thuộc danh mục dịch vụ Thoại (rớt cuộc gọi, không gọi đi được, lỗi hiển thị số), Tin nhắn (không nhận mã OTP ngân hàng) hoặc Chuyển mạng giữ số.",
        "Hệ sinh thái tự động đóng hiện tại chuyên trách sâu cho Mobile Internet (BTools/SAPC Data). Các dịch vụ Thoại/SMS cần đối soát CDR và Core MSC riêng biệt.",
        "Đóng nhầm nội dung data cho phiếu sự cố thoại/tin nhắn.",
        "KTV chuyển sang tab nghiệp vụ Thoại/SMS để tra cứu CDR cuộc gọi và phân tích nguyên nhân lỗi riêng theo quy trình thoại.",
        "Kiểm tra tính năng VoLTE trên máy và HSS; đối soát mã lỗi SIP cuộc gọi trên CDR."
    ),
    (
        "MAN_07", "Dữ liệu BTools chưa đồng bộ hoặc rỗng không đủ cơ sở kết luận",
        "BTools đang trong thời gian bảo dưỡng hoặc dữ liệu chưa đẩy về kịp, toàn bộ bảng phiên trắng trơn; HSS và CEM chưa có dữ liệu đồng bộ.",
        "Không đủ dữ liệu chứng cứ kỹ thuật để khẳng định nguyên nhân lỗi.",
        "Đưa ra kết luận đóng phiếu sai thực tế, thiếu tính thuyết phục.",
        "KTV bấm 'Quét lại' (Refresh) để kéo lại dữ liệu mới nhất từ BTools; nếu vẫn thiếu thì tra cứu trực tiếp trên hệ thống BTools web.",
        "Kiểm tra lại sau 15 phút hoặc liên hệ khách hàng xác minh thêm thông tin."
    )
]

for row_idx, r_data in enumerate(data_manual, start=4):
    stt = row_idx - 3
    row_vals = [stt] + list(r_data)
    for col_idx, val in enumerate(row_vals, start=1):
        cell = ws_manual.cell(row=row_idx, column=col_idx, value=val)
        cell.border = border_cell
        cell.font = font_body
        if col_idx in (1, 2):
            cell.alignment = align_top_center
            if col_idx == 2:
                cell.font = font_code
        else:
            cell.alignment = align_top_left
            if col_idx == 3:
                cell.font = font_body_bold
                cell.fill = PatternFill(start_color=COLOR_MANUAL_BG, end_color=COLOR_MANUAL_BG, fill_type="solid")
            elif col_idx == 7:
                cell.font = font_body_bold
    ws_manual.row_dimensions[row_idx].height = 54
    if row_idx % 2 == 1:
        for col_idx in range(1, len(row_vals) + 1):
            if col_idx != 3:
                ws_manual.cell(row=row_idx, column=col_idx).fill = fill_zebra

auto_adjust_columns(ws_manual, len(headers_manual))

# Luu file Excel
wb.save(OUTPUT_FILE)
print(f"Xuat thanh cong file Excel: {OUTPUT_FILE}")
